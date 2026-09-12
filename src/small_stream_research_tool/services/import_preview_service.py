"""매핑된 원본 행의 identity Preview. 판정/표시는 DB Import·영구 QC가 아니다."""

from collections import Counter
from collections.abc import Iterator

from openpyxl.utils import get_column_letter

from small_stream_research_tool.models.column_mapping import MappingStatus
from small_stream_research_tool.models.excel import ExcelCell, ExcelRow
from small_stream_research_tool.models.import_preview import (
    ImportPreviewIssue,
    ImportPreviewResult,
    ImportPreviewRow,
    ImportPreviewSummary,
    PreviewDisplayValue,
    PreviewFieldPolicy,
    PreviewMappedValue,
    PreviewStatus,
)
from small_stream_research_tool.models.import_preview_errors import (
    InvalidPreviewArgumentError,
    PreviewSourceChangedError,
)
from small_stream_research_tool.models.stream_code import CodeStatus, ComparisonStatus
from small_stream_research_tool.repositories.stream_lookup_repository import StreamLookupRepository
from small_stream_research_tool.services.column_mapping_service import ColumnMappingService
from small_stream_research_tool.services.dictionary_service import DictionaryService
from small_stream_research_tool.services.excel_reader import ExcelReader
from small_stream_research_tool.services.stream_code_service import (
    COMPONENT_WIDTHS,
    validate_stream_code_cells,
)
from small_stream_research_tool.utils.file_hash import file_sha256

IDENTITY_NAMES = frozenset((*COMPONENT_WIDTHS, "stream_code"))
MASKED_VALUE = "[MASKED]"


def _validate_policy(policy):
    if type(policy) is not PreviewFieldPolicy:
        raise InvalidPreviewArgumentError("PreviewFieldPolicy가 필요합니다.")
    for names in (policy.excluded_internal_names, policy.masked_internal_names):
        if type(names) is not frozenset or any(
            not isinstance(name, str) or not name or name != name.strip() for name in names
        ):
            raise InvalidPreviewArgumentError("정확한 internal_name의 frozenset이 필요합니다.")
        if names & IDENTITY_NAMES:
            raise InvalidPreviewArgumentError("연구 식별자인 관리코드 필드는 숨김 대상이 아닙니다.")


def _row_cells(row):
    if (
        type(row) is not ExcelRow
        or type(row.row_index) is not int
        or not 1 <= row.row_index <= 1048576
    ):
        raise InvalidPreviewArgumentError("1-based 원본 행 번호의 ExcelRow가 필요합니다.")
    result = {}
    for cell in row.cells:
        if (
            type(cell) is not ExcelCell
            or cell.row_index != row.row_index
            or (type(cell.column_index) is not int or not 1 <= cell.column_index <= 16384)
        ):
            raise InvalidPreviewArgumentError("셀과 원본 행/열 위치가 올바르지 않습니다.")
        if cell.column_index in result or cell.column_letter != get_column_letter(
            cell.column_index
        ):
            raise InvalidPreviewArgumentError("원본 열 번호가 중복되거나 열 문자와 다릅니다.")
        result[cell.column_index] = cell
    return result


def preview_summary(rows) -> ImportPreviewSummary:
    counts = Counter(row.status for row in rows)
    return ImportPreviewSummary(sum(counts.values()), *(counts[status] for status in PreviewStatus))


class ImportPreviewService:
    def __init__(
        self, dictionary_service: DictionaryService, stream_lookup: StreamLookupRepository
    ):
        self._dictionary = dictionary_service
        self._mapping = ColumnMappingService(dictionary_service)
        self._streams = stream_lookup

    def _prepare(self, mappings):
        mappings = self._mapping.revalidate_mappings(mappings)
        selected, issues = [], []
        for mapping in mappings:
            if mapping.mapping_status == MappingStatus.NEEDS_REVIEW:
                issues.append(
                    ImportPreviewIssue(
                        "MAPPING_NEEDS_REVIEW",
                        "매핑을 다시 확인해야 합니다.",
                        True,
                        mapping.source_column_index,
                    )
                )
            elif mapping.mapping_status in (MappingStatus.AUTO_MAPPED, MappingStatus.USER_MAPPED):
                item = self._dictionary.get_item(mapping.dictionary_id)
                if item is None:
                    issues.append(
                        ImportPreviewIssue(
                            "MAPPING_NEEDS_REVIEW",
                            "매핑 대상을 확인할 수 없습니다.",
                            True,
                            mapping.source_column_index,
                        )
                    )
                else:
                    selected.append((mapping, item))
        names = Counter(item.internal_name for _, item in selected)
        duplicates = {name for name, count in names.items() if count > 1}
        for mapping, item in selected:
            if item.internal_name in duplicates:
                issues.append(
                    ImportPreviewIssue(
                        "DUPLICATE_SEMANTIC_MAPPING",
                        "동일 항목에 여러 원본 컬럼이 연결되어 자동 선택할 수 없습니다.",
                        True,
                        mapping.source_column_index,
                    )
                )
        available = {item.internal_name for _, item in selected} - duplicates
        if "stream_code" not in available and not set(COMPONENT_WIDTHS) <= available:
            issues.append(
                ImportPreviewIssue(
                    "IDENTITY_MAPPING_INCOMPLETE",
                    "전체 관리코드 또는 네 구성요소를 사용할 수 있도록 매핑해야 합니다.",
                    True,
                )
            )
        return tuple(selected), tuple(issues), duplicates

    def iter_preview_rows(
        self,
        rows,
        mappings,
        *,
        excluded_rows=frozenset(),
        field_policy=None,
        include_blank=False,
    ) -> Iterator[ImportPreviewRow]:
        field_policy = PreviewFieldPolicy() if field_policy is None else field_policy
        _validate_policy(field_policy)
        if type(include_blank) is not bool:
            raise InvalidPreviewArgumentError("빈 행 포함 여부는 bool이어야 합니다.")
        try:
            excluded = frozenset(excluded_rows)
        except TypeError:
            raise InvalidPreviewArgumentError("제외할 원본 행 번호 집합이 필요합니다.") from None
        if any(type(row) is not int or not 1 <= row <= 1048576 for row in excluded):
            raise InvalidPreviewArgumentError("제외 행은 1-based 원본 행 번호여야 합니다.")
        selected, issues, duplicates = self._prepare(tuple(mappings))
        return self._iter(rows, selected, issues, duplicates, excluded, field_policy, include_blank)

    def _iter(self, rows, selected, initial_issues, duplicates, excluded, policy, include_blank):
        seen_rows = set()
        for row in rows:
            cells = _row_cells(row)
            if row.row_index in seen_rows:
                raise InvalidPreviewArgumentError("Preview 입력에 원본 행 번호가 중복되었습니다.")
            seen_rows.add(row.row_index)
            if row.is_blank and not include_blank and row.row_index not in excluded:
                continue
            yield self._preview_row(
                row, cells, selected, initial_issues, duplicates, row.row_index in excluded, policy
            )

    def _preview_row(self, row, cells, selected, initial_issues, duplicates, excluded, policy):
        issues, mapped, display, code_cells = list(initial_issues), [], [], {}
        stream_name = None
        for mapping, item in selected:
            index, name = mapping.source_column_index, item.internal_name
            cell = cells.get(index)
            if cell is None:
                issues.append(
                    ImportPreviewIssue(
                        "SOURCE_COLUMN_MISSING",
                        "매핑에 해당하는 원본 셀 위치가 없습니다.",
                        True,
                        index,
                    )
                )
                continue
            mapped.append(PreviewMappedValue(index, item.dictionary_id, name, cell))
            if name in IDENTITY_NAMES and name not in duplicates:
                code_cells[name] = cell
            if name not in policy.excluded_internal_names:
                value = MASKED_VALUE if name in policy.masked_internal_names else cell.value
                display.append(PreviewDisplayValue(index, name, value))
                if (
                    name == "stream_name"
                    and name not in duplicates
                    and isinstance(value, str)
                    and (not cell.is_formula and cell.value_type not in ("f", "e"))
                ):
                    stream_name = value
        validation = validate_stream_code_cells(code_cells)
        for part in (validation.source, *validation.components):
            if part.status != CodeStatus.VALID:
                blocking = part.status != CodeStatus.MISSING
                issues.append(
                    ImportPreviewIssue(
                        part.code,
                        part.message,
                        blocking,
                        code_cells[part.name].column_index if part.name in code_cells else None,
                    )
                )
        effective = None
        if validation.comparison_status == ComparisonStatus.MISMATCH:
            issues.append(
                ImportPreviewIssue(
                    "SOURCE_COMPONENT_MISMATCH",
                    "원본 관리코드와 구성요소 조합값이 일치하지 않습니다.",
                    True,
                )
            )
        elif validation.source_valid:
            # 구성요소 누락은 원본 코드 후보를 막지 않는다. 잘못된/모호한 값은 별도 차단한다.
            effective = validation.normalized_source_code
        elif validation.source.status == CodeStatus.MISSING and validation.generated_valid:
            effective = validation.generated_code
            issues.append(
                ImportPreviewIssue(
                    "SOURCE_MISSING_GENERATED_AVAILABLE",
                    "원본 코드는 없으며 구성요소로 생성한 후보를 사용합니다.",
                    False,
                )
            )
        if any(
            part.status
            in (CodeStatus.INVALID_FORMAT, CodeStatus.INVALID_LENGTH, CodeStatus.AMBIGUOUS)
            for part in (validation.source, *validation.components)
        ):
            effective = None
        if effective is None:
            issues.append(
                ImportPreviewIssue(
                    "IDENTITY_UNAVAILABLE", "안전하게 사용할 관리코드를 확보할 수 없습니다.", True
                )
            )
        found = None
        if excluded:
            status = PreviewStatus.EXCLUDED
        elif any(issue.blocking for issue in issues):
            status = PreviewStatus.NEEDS_REVIEW
        else:
            found = self._streams.exists_by_stream_code(effective)
            status = PreviewStatus.EXISTING_STREAM if found else PreviewStatus.NEW_STREAM
        return ImportPreviewRow(
            row.row_index,
            status,
            effective,
            stream_name,
            validation,
            tuple(mapped),
            tuple(display),
            tuple(issues),
            excluded,
            found,
        )

    def build_preview(self, rows, mappings, **options) -> ImportPreviewResult:
        """호출자가 정한 입력 범위의 결과를 모은다. 대량 순회는 iter_preview_rows를 사용한다."""
        result = tuple(self.iter_preview_rows(rows, mappings, **options))
        return ImportPreviewResult(result, preview_summary(result))

    def rebuild_from_workspace(self, workspace_service, current_user_id, **options):
        resumed = self._mapping.resume_workspace(workspace_service, current_user_id)
        if resumed is None:
            return None
        draft = resumed.resumed
        if draft.selected_sheet_name is None or draft.header_start_row is None:
            raise InvalidPreviewArgumentError("Preview 재생성 전에 시트·헤더를 선택해야 합니다.")
        with ExcelReader(draft.source_file_path) as reader:
            rows = reader.iter_rows(
                draft.selected_sheet_name,
                header_start_row=draft.header_start_row,
                header_end_row=draft.header_end_row,
                data_start_row=draft.data_start_row,
                include_blank=True,
            )
            preview = self.build_preview(rows, draft.column_mappings, **options)
        try:
            digest = file_sha256(draft.source_file_path)
        except OSError:
            raise PreviewSourceChangedError(
                "Preview 재생성 후 원본을 확인할 수 없습니다."
            ) from None
        if digest != draft.source_file_sha256:
            raise PreviewSourceChangedError("Preview 재생성 중 원본이 변경되었습니다.")
        return preview
