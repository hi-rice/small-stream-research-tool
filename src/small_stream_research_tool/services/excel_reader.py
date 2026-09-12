"""원본 저장·DB 접근·의미 해석 없이 .xlsx 구조와 native 값을 읽는다."""

from collections.abc import Iterator
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.worksheet.worksheet import Worksheet

from small_stream_research_tool.models.excel import (
    ExcelCell,
    ExcelColumn,
    ExcelFormula,
    ExcelHeaderPart,
    ExcelPreview,
    ExcelRow,
    ExcelSheetInfo,
    ExcelWorkbookInfo,
)
from small_stream_research_tool.models.excel_errors import (
    ExcelFileNotFoundError,
    ExcelReaderClosedError,
    ExcelSheetNotFoundError,
    ExcelWorkbookReadError,
    InvalidDataRangeError,
    InvalidHeaderRangeError,
    UnsupportedExcelFormatError,
    UnsupportedExcelSheetError,
)

HEADER_SEPARATOR = " | "


def _positive_integer(value: object) -> bool:
    return type(value) is int and value > 0


def _cell(sheet: Worksheet, row: int, column: int) -> ExcelCell:
    source = sheet.cell(row, column)
    value = source.value
    if isinstance(value, (ArrayFormula, DataTableFormula)):
        value = ExcelFormula(
            kind=value.t,
            expression=value.text if isinstance(value, ArrayFormula) else None,
            attributes=tuple(dict(value).items()),
        )
    return ExcelCell(
        row,
        column,
        get_column_letter(column),
        value,
        source.data_type,
        source.data_type == "f",
        source.number_format,
    )


class ExcelReader:
    """한 번 로딩하며 context manager로 사용한다. 닫힌 Reader는 재사용하지 않는다."""

    def __init__(self, path: str | Path):
        self._workbook = None
        try:
            self._path = Path(path).absolute()
            if not self._path.exists():
                raise ExcelFileNotFoundError("Excel 파일이 존재하지 않습니다.")
            if not self._path.is_file():
                raise ExcelWorkbookReadError("일반 파일을 선택해야 합니다.")
            if self._path.suffix.lower() != ".xlsx":
                raise UnsupportedExcelFormatError(".xlsx 형식만 지원합니다.")
        except (OSError, TypeError, ValueError):
            raise ExcelWorkbookReadError("Excel 파일 경로에 접근할 수 없습니다.") from None
        try:
            # 일반 모드가 병합 메타데이터를 제공한다. 수식은 계산/캐시 대체하지 않는다.
            # 읽기 전용 파일 handle을 직접 소유해 로딩 실패 때에도 닫는다.
            with self._path.open("rb") as source:
                self._workbook = load_workbook(
                    source,
                    read_only=False,
                    data_only=False,
                    keep_vba=False,
                    keep_links=False,
                    rich_text=False,
                )
        except Exception:
            raise ExcelWorkbookReadError("Excel workbook을 읽을 수 없습니다.") from None

    def _ensure_open(self) -> None:
        if self._workbook is None:
            raise ExcelReaderClosedError("Excel Reader가 닫혀 있습니다.")

    def __enter__(self) -> "ExcelReader":
        self._ensure_open()
        return self

    def __exit__(self, exc_type, exc_value, traceback) -> None:
        self.close()

    def close(self) -> None:
        if self._workbook is not None:
            workbook, self._workbook = self._workbook, None
            workbook.close()

    @property
    def workbook_info(self) -> ExcelWorkbookInfo:
        self._ensure_open()
        return ExcelWorkbookInfo(self._path, tuple(self._workbook.sheetnames))

    def sheet_info(self, sheet_name: str) -> ExcelSheetInfo:
        self._ensure_open()
        if sheet_name not in self._workbook.sheetnames:
            raise ExcelSheetNotFoundError("요청한 시트가 없습니다.")
        sheet = self._workbook[sheet_name]
        is_worksheet = isinstance(sheet, Worksheet)
        return ExcelSheetInfo(
            sheet.title,
            self._workbook.sheetnames.index(sheet_name) + 1,
            sheet.max_row if is_worksheet else None,
            sheet.max_column if is_worksheet else None,
            tuple(sorted(str(r) for r in sheet.merged_cells.ranges)) if is_worksheet else (),
            sheet.sheet_state,
            "worksheet" if is_worksheet else "chartsheet",
        )

    def _sheet(self, sheet_name: str) -> Worksheet:
        if self.sheet_info(sheet_name).kind != "worksheet":
            raise UnsupportedExcelSheetError("셀 데이터가 있는 worksheet를 선택해야 합니다.")
        return self._workbook[sheet_name]

    @staticmethod
    def _validate_header(sheet: Worksheet, start: int, end: int) -> None:
        if not (
            _positive_integer(start) and _positive_integer(end) and start <= end <= sheet.max_row
        ):
            raise InvalidHeaderRangeError(
                "헤더 범위는 시트 내부의 1-based 시작·종료 행이어야 합니다."
            )

    def read_columns(
        self,
        sheet_name: str,
        *,
        header_start_row: int,
        header_end_row: int,
    ) -> tuple[ExcelColumn, ...]:
        sheet = self._sheet(sheet_name)
        self._validate_header(sheet, header_start_row, header_end_row)
        ranges = tuple(sheet.merged_cells.ranges)
        columns = []
        for column in range(1, sheet.max_column + 1):
            parts = []
            display = []
            seen_anchors = set()
            for row in range(header_start_row, header_end_row + 1):
                original = _cell(sheet, row, column)
                merged = next(
                    (
                        r
                        for r in ranges
                        if r.min_row <= row <= r.max_row and r.min_col <= column <= r.max_col
                    ),
                    None,
                )
                anchor = _cell(sheet, merged.min_row, merged.min_col) if merged else None
                part = ExcelHeaderPart(original, anchor, str(merged) if merged else None)
                parts.append(part)
                # 동일한 세로 병합 anchor의 표시 중복만 생략한다. 원본 parts는 모두 유지한다.
                if anchor is not None and anchor.coordinate in seen_anchors:
                    continue
                if anchor is not None:
                    seen_anchors.add(anchor.coordinate)
                value = part.effective_value
                if value is not None:
                    display.append(str(value))
            columns.append(
                ExcelColumn(
                    column, get_column_letter(column), tuple(parts), HEADER_SEPARATOR.join(display)
                )
            )
        return tuple(columns)

    def iter_rows(
        self,
        sheet_name: str,
        *,
        header_start_row: int,
        header_end_row: int,
        data_start_row: int | None = None,
        include_blank: bool = True,
    ) -> Iterator[ExcelRow]:
        sheet = self._sheet(sheet_name)
        self._validate_header(sheet, header_start_row, header_end_row)
        start = header_end_row + 1 if data_start_row is None else data_start_row
        if not (_positive_integer(start) and header_end_row < start <= sheet.max_row + 1):
            raise InvalidDataRangeError(
                "데이터 시작 행은 헤더 다음부터 시트 마지막 행+1까지입니다."
            )
        if type(include_blank) is not bool:
            raise InvalidDataRangeError("빈 행 포함 여부는 bool이어야 합니다.")
        return self._iter_rows(sheet, start, sheet.max_row, include_blank)

    def _iter_rows(
        self,
        sheet: Worksheet,
        start: int,
        end: int,
        include_blank: bool,
    ) -> Iterator[ExcelRow]:
        width = sheet.max_column
        for row_index in range(start, end + 1):
            self._ensure_open()
            row = ExcelRow(
                row_index, tuple(_cell(sheet, row_index, c) for c in range(1, width + 1))
            )
            if include_blank or not row.is_blank:
                yield row

    def read_preview(
        self,
        sheet_name: str,
        *,
        header_start_row: int,
        header_end_row: int,
        start_row: int | None = None,
        row_count: int = 10,
    ) -> ExcelPreview:
        """지정한 실제 행 범위를 제한해서 반환한다. 빈 행도 row_count에 포함한다."""
        columns = self.read_columns(
            sheet_name, header_start_row=header_start_row, header_end_row=header_end_row
        )
        sheet = self._sheet(sheet_name)
        start = header_end_row + 1 if start_row is None else start_row
        if not (_positive_integer(start) and header_end_row < start <= sheet.max_row + 1):
            raise InvalidDataRangeError(
                "Preview 시작 행은 헤더 다음부터 시트 마지막 행+1까지입니다."
            )
        if type(row_count) is not int or row_count < 0:
            raise InvalidDataRangeError("Preview 행 수는 0 이상의 정수여야 합니다.")
        rows = tuple(self._iter_rows(sheet, start, min(sheet.max_row, start + row_count - 1), True))
        return ExcelPreview(self.sheet_info(sheet_name), columns, rows)
