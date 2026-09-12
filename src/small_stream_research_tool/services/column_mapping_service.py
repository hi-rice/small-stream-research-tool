"""등록 별칭만 조회한다. 사용자 선택·원본 metadata를 덮어쓰거나 DB에 저장하지 않는다."""

from dataclasses import replace

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange

from small_stream_research_tool.models.column_mapping import (
    ColumnMappingDraft,
    MappingHeaderPart,
    MappingMethod,
    MappingStatus,
    MappingSummary,
)
from small_stream_research_tool.models.dictionary_errors import InvalidDictionaryDefinitionError
from small_stream_research_tool.models.excel import ExcelColumn
from small_stream_research_tool.models.mapping_errors import (
    InvalidMappingError,
    MappingTargetUnavailableError,
)
from small_stream_research_tool.models.workspace import WorkspaceResumeResult
from small_stream_research_tool.services.dictionary_service import DictionaryService


def _positive(value):
    return type(value) is int and 0 < value <= 9223372036854775807


def validate_mapping(mapping: ColumnMappingDraft) -> None:
    """DB 조회 없이 metadata와 상태/방법 조합만 검증한다."""
    if type(mapping) is not ColumnMappingDraft:
        raise InvalidMappingError("ColumnMappingDraft가 필요합니다.")
    if not (_positive(mapping.source_column_index) and mapping.source_column_index <= 16384):
        raise InvalidMappingError("원본 열 번호는 Excel의 1-based 범위여야 합니다.")
    if mapping.source_column_letter != get_column_letter(mapping.source_column_index):
        raise InvalidMappingError("원본 열 번호와 열 문자가 다릅니다.")
    if not isinstance(mapping.source_header, str) or "\0" in mapping.source_header:
        raise InvalidMappingError("원본 헤더 문자열이 필요합니다.")
    if mapping.source_scope is not None and (
        not isinstance(mapping.source_scope, str)
        or not mapping.source_scope.strip()
        or "\0" in mapping.source_scope
    ):
        raise InvalidMappingError("source_scope 형식이 올바르지 않습니다.")
    if type(mapping.source_header_parts) is not tuple:
        raise InvalidMappingError("헤더 parts는 tuple이어야 합니다.")
    previous_row = 0
    for part in mapping.source_header_parts:
        if type(part) is not MappingHeaderPart or not _positive(part.row_index):
            raise InvalidMappingError("헤더 part 형식이 올바르지 않습니다.")
        if not previous_row < part.row_index <= 1048576:
            raise InvalidMappingError("헤더 행은 원본 순서여야 합니다.")
        previous_row = part.row_index
        for value in (part.text, part.anchor_text, part.merged_range):
            if value is not None and (not isinstance(value, str) or "\0" in value):
                raise InvalidMappingError("헤더 metadata는 문자열 또는 None이어야 합니다.")
        if part.merged_range is None:
            if any(v is not None for v in (part.anchor_row, part.anchor_column, part.anchor_text)):
                raise InvalidMappingError("병합 범위 없는 anchor는 허용하지 않습니다.")
        elif not (
            _positive(part.anchor_row)
            and _positive(part.anchor_column)
            and part.anchor_row <= 1048576
            and part.anchor_column <= 16384
        ):
            raise InvalidMappingError("병합 anchor 위치가 올바르지 않습니다.")
        else:
            try:
                merged = CellRange(part.merged_range)
                valid_range = (
                    merged.title is None
                    and merged.min_row == part.anchor_row
                    and merged.min_col == part.anchor_column
                    and merged.min_row <= part.row_index <= merged.max_row
                    and merged.min_col <= mapping.source_column_index <= merged.max_col
                )
            except (ValueError, TypeError):
                valid_range = False
            if not valid_range:
                raise InvalidMappingError("병합 범위와 원본/anchor 위치가 다릅니다.")
    if not isinstance(mapping.mapping_status, MappingStatus) or not isinstance(
        mapping.mapping_method, MappingMethod
    ):
        raise InvalidMappingError("지원하는 매핑 상태·방법이 필요합니다.")
    expected_methods = {
        MappingStatus.AUTO_MAPPED: (MappingMethod.AUTO_ALIAS,),
        MappingStatus.USER_MAPPED: (MappingMethod.USER,),
        MappingStatus.UNMAPPED: (MappingMethod.NONE,),
        MappingStatus.DO_NOT_MAP: (MappingMethod.USER,),
        MappingStatus.NEEDS_REVIEW: (MappingMethod.AUTO_ALIAS, MappingMethod.USER),
    }
    if mapping.mapping_method not in expected_methods[mapping.mapping_status]:
        raise InvalidMappingError("매핑 상태·방법 조합이 올바르지 않습니다.")
    if mapping.mapping_status in (MappingStatus.UNMAPPED, MappingStatus.DO_NOT_MAP):
        if mapping.dictionary_id is not None:
            raise InvalidMappingError("미매핑/제외 상태에는 사전 ID를 지정하지 않습니다.")
    elif not _positive(mapping.dictionary_id):
        raise InvalidMappingError("매핑/재검토 상태에는 양의 사전 ID가 필요합니다.")


def validate_mappings(mappings):
    indexes = set()
    for mapping in mappings:
        validate_mapping(mapping)
        if mapping.source_column_index in indexes:
            raise InvalidMappingError("같은 원본 열 번호의 매핑이 중복되었습니다.")
        indexes.add(mapping.source_column_index)


def mapping_summary(mappings) -> MappingSummary:
    mappings = tuple(mappings)
    validate_mappings(mappings)
    return MappingSummary(
        len(mappings),
        *(sum(m.mapping_status == status for m in mappings) for status in MappingStatus),
    )


class ColumnMappingService:
    def __init__(self, dictionary_service: DictionaryService):
        self._dictionary = dictionary_service

    def _target_usable(self, dictionary_id):
        item = self._dictionary.get_item(dictionary_id)
        if item is None or not item.is_active or item.deprecated_version_id is not None:
            return False
        category = self._dictionary.get_category(item.category_id)
        unit = self._dictionary.get_unit(item.unit_id) if item.unit_id is not None else None
        return (
            category is not None
            and category.is_active
            and (item.unit_id is None or (unit is not None and unit.is_active))
        )

    def _auto(self, mapping):
        validate_mapping(mapping)
        target = None
        if mapping.source_header.strip():
            try:
                target = self._dictionary.find_dictionary_by_header(
                    mapping.source_header, mapping.source_scope
                )
            except InvalidDictionaryDefinitionError:
                raise InvalidMappingError("별칭 조회 입력이 올바르지 않습니다.") from None
        return replace(
            mapping,
            dictionary_id=target.dictionary_id if target else None,
            mapping_method=MappingMethod.AUTO_ALIAS if target else MappingMethod.NONE,
            mapping_status=MappingStatus.AUTO_MAPPED if target else MappingStatus.UNMAPPED,
        )

    def build_initial_mappings(self, columns, source_scope=None):
        mappings = []
        for column in columns:
            if type(column) is not ExcelColumn:
                raise InvalidMappingError("ExcelColumn 목록이 필요합니다.")
            if any(part.cell.column_index != column.column_index for part in column.header_parts):
                raise InvalidMappingError("헤더 part와 원본 열 번호가 다릅니다.")
            parts = tuple(
                MappingHeaderPart(
                    part.cell.row_index,
                    None if part.cell.value is None else str(part.cell.value),
                    part.merged_anchor.row_index if part.merged_anchor else None,
                    part.merged_anchor.column_index if part.merged_anchor else None,
                    (None if part.merged_anchor.value is None else str(part.merged_anchor.value))
                    if part.merged_anchor
                    else None,
                    part.merged_range,
                )
                for part in column.header_parts
            )
            mapping = ColumnMappingDraft(
                column.column_index,
                column.column_letter,
                column.display_header,
                parts,
                source_scope=source_scope,
            )
            mappings.append(self._auto(mapping))
        validate_mappings(mappings)
        return tuple(mappings)

    def set_user_mapping(self, mapping, dictionary_id):
        validate_mapping(mapping)
        if not _positive(dictionary_id):
            raise InvalidMappingError("양의 사전 ID가 필요합니다.")
        if not self._target_usable(dictionary_id):
            raise MappingTargetUnavailableError("사용 가능한 활성 사전 항목이 필요합니다.")
        return replace(
            mapping,
            dictionary_id=dictionary_id,
            mapping_method=MappingMethod.USER,
            mapping_status=MappingStatus.USER_MAPPED,
        )

    def clear_mapping(self, mapping):
        validate_mapping(mapping)
        return replace(
            mapping,
            dictionary_id=None,
            mapping_method=MappingMethod.NONE,
            mapping_status=MappingStatus.UNMAPPED,
        )

    def mark_do_not_map(self, mapping):
        validate_mapping(mapping)
        return replace(
            mapping,
            dictionary_id=None,
            mapping_method=MappingMethod.USER,
            mapping_status=MappingStatus.DO_NOT_MAP,
        )

    def refresh_auto_mappings(self, mappings):
        mappings = tuple(mappings)
        validate_mappings(mappings)
        return tuple(
            self._auto(m)
            if m.mapping_status in (MappingStatus.AUTO_MAPPED, MappingStatus.UNMAPPED)
            else m
            for m in mappings
        )

    def revalidate_mappings(self, mappings):
        mappings = tuple(mappings)
        validate_mappings(mappings)
        result = []
        for mapping in mappings:
            needs_review = mapping.dictionary_id is not None and not self._target_usable(
                mapping.dictionary_id
            )
            if mapping.mapping_status == MappingStatus.AUTO_MAPPED and not needs_review:
                # alias 비활성/재연결도 감지하되 새 target으로 자동 교체하지 않는다.
                current = self._auto(mapping)
                needs_review = current.dictionary_id != mapping.dictionary_id
            result.append(
                replace(mapping, mapping_status=MappingStatus.NEEDS_REVIEW)
                if needs_review
                else mapping
            )
        return tuple(result)

    def resume_workspace(self, workspace_service, current_user_id):
        original = workspace_service.load_workspace(current_user_id)
        if original is None:
            return None
        resumed = replace(
            original, column_mappings=self.revalidate_mappings(original.column_mappings)
        )
        return WorkspaceResumeResult(original, resumed)
