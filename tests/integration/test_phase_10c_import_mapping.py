"""합성 원본과 임시 DB에서 10B→10C 매핑·Preview 경계를 검증한다."""

import os
import time
from contextlib import closing

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from openpyxl import Workbook
from PySide6.QtWidgets import QApplication

from small_stream_research_tool.app.controller import ApplicationController
from small_stream_research_tool.database import connect_database, initialize_database
from small_stream_research_tool.models.column_mapping import MappingStatus
from small_stream_research_tool.models.import_mapping_workflow import MappingWorkflowError
from small_stream_research_tool.repositories.dictionary_repository import DictionaryRepository
from small_stream_research_tool.services.dictionary_service import DictionaryService
from small_stream_research_tool.services.import_core_dictionary_bootstrap import (
    ImportCoreDictionaryBootstrapService,
)
from small_stream_research_tool.services.import_inspection_service import ImportInspectionService
from small_stream_research_tool.services.import_mapping_workflow_service import (
    ImportMappingWorkflowService,
)
from small_stream_research_tool.services.research_dictionary_bootstrap import (
    ResearchDictionaryBootstrapService,
)
from small_stream_research_tool.ui.import_mapping import ImportMappingView

CORE = (
    ("stream_code", "관리코드"),
    ("province_code", "시도코드"),
    ("city_county_code", "시군구코드"),
    ("town_code", "읍면동코드"),
    ("stream_serial_no", "일련번호"),
    ("stream_name", "소하천명"),
)


def _snapshot(path):
    with closing(connect_database(path)) as connection:
        return tuple(connection.iterdump())


def _bootstrap(path):
    with closing(connect_database(path)) as connection:
        ResearchDictionaryBootstrapService(connection).bootstrap()


def _workbook(path, *, invalid=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "합성 자료"
    sheet.append([*(label for _name, label in CORE), "유역면적", "Service Key", "미확정 열"])
    sheet.append(["01234567001", "01", "234", "567", "001", "합성 하천", 2.5, "가상값", "A"])
    if invalid:
        sheet.append([1234567002, "01", "234", "567", "002", "다른 합성", 3.5, "가상값", "B"])
    book.save(path)
    book.close()


def _identity_workbook(path, headers, values):
    book = Workbook()
    sheet = book.active
    sheet.title = "합성 자료"
    sheet.append(headers)
    sheet.append(values)
    book.save(path)
    book.close()


def _source(path, db, workspace):
    inspection = ImportInspectionService(db, workspace)
    state = inspection.inspect(path)
    structured = inspection.structure(state, "합성 자료", 1, 1, 2)
    return inspection.save(structured, 1)


def _wait(app, predicate):
    limit = time.monotonic() + 10
    while time.monotonic() < limit:
        app.processEvents()
        if predicate():
            return
        time.sleep(0.02)
    pytest.fail("합성 GUI 작업이 끝나지 않았습니다.")


def test_dictionary_required_and_mapping_checkpoint(tmp_path):
    db = tmp_path / "research.sqlite3"
    workbook = tmp_path / "합성.xlsx"
    workspace = tmp_path / "workspace"
    initialize_database(db)
    _workbook(workbook)
    source = _source(workbook, db, workspace)
    workflow = ImportMappingWorkflowService(db, workspace)
    with pytest.raises(MappingWorkflowError, match="연구 사전"):
        workflow.start(source, 1)
    _bootstrap(db)
    before = _snapshot(db)
    state = workflow.start(source, 1)
    assert state.mappings[0].mapping_status == MappingStatus.UNMAPPED
    state = workflow.set_scope(state, "NATIONAL_2024", 1)
    assert state.mappings[0].mapping_status == MappingStatus.AUTO_MAPPED
    assert state.mappings[6].mapping_status == MappingStatus.AUTO_MAPPED
    assert state.mappings[7].mapping_status == MappingStatus.DO_NOT_MAP
    assert state.mappings[8].mapping_status == MappingStatus.UNMAPPED
    assert state.rows[7].header == "제외 대상 열"
    assert "가상값" not in repr(state)
    assert state.saved_at is not None
    metadata = (workspace / "user_1.json").read_text(encoding="utf-8")
    assert "Service Key" not in metadata and "가상값" not in metadata
    assert _snapshot(db) == before
    basin = state.mappings[6].dictionary_id
    state = workflow.update(state, 9, "map", basin, 1)
    assert state.mappings[8].mapping_status == MappingStatus.USER_MAPPED
    state = workflow.update(state, 9, "exclude", None, 1)
    assert state.mappings[8].mapping_status == MappingStatus.DO_NOT_MAP
    state = workflow.update(state, 9, "clear", None, 1)
    assert state.mappings[8].mapping_status == MappingStatus.UNMAPPED
    state = workflow.update(state, 9, "exclude", None, 1)
    assert _snapshot(db) == before
    resumed = workflow.start(state.source, 1)
    assert resumed.mappings == state.mappings


def test_production_bootstrap_prepares_stream_identity(tmp_path):
    db = tmp_path / "research.sqlite3"
    workbook = tmp_path / "합성.xlsx"
    workspace = tmp_path / "workspace"
    initialize_database(db)
    with closing(connect_database(db)) as connection:
        ResearchDictionaryBootstrapService(connection).bootstrap()
    _workbook(workbook)
    source = _source(workbook, db, workspace)
    before = _snapshot(db)

    workflow = ImportMappingWorkflowService(db, workspace)
    state = workflow.set_scope(workflow.start(source, 1), "NATIONAL_2024", 1)
    with closing(connect_database(db)) as connection:
        names = {item.internal_name for item in DictionaryRepository(connection).list_items()}
    assert {name for name, _label in CORE} <= names
    assert all(
        mapping.mapping_status == MappingStatus.AUTO_MAPPED for mapping in state.mappings[:7]
    )

    result = workflow.preview(state, 1)
    assert result.preview.total == 1
    assert result.preview.ready == 1
    assert result.preview.blocked == 0
    assert result.preview.ready_for_import_preparation
    assert _snapshot(db) == before


def test_direct_code_maps_as_text_for_existing_stream(tmp_path):
    db = tmp_path / "research.sqlite3"
    workbook = tmp_path / "직접코드.xlsx"
    workspace = tmp_path / "workspace"
    initialize_database(db)
    _bootstrap(db)
    with closing(connect_database(db)) as connection:
        connection.execute(
            "INSERT INTO small_stream (stream_code,province_code,city_county_code,town_code,"
            "stream_serial_no,stream_name,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?)",
            (
                "01234567001",
                "01",
                "234",
                "567",
                "001",
                "기존 합성 하천",
                "2026-09-23T00:00:00Z",
                "2026-09-23T00:00:00Z",
            ),
        )
    _identity_workbook(workbook, ["관리코드", "유역면적"], ["01234567001", 2.5])
    source = _source(workbook, db, workspace)
    workflow = ImportMappingWorkflowService(db, workspace)
    state = workflow.set_scope(workflow.start(source, 1), "NATIONAL_2024", 1)
    result = workflow.preview(state, 1)
    assert result.preview.ready_for_import_preparation
    assert result.preview.displayed[0].stream_code == "01234567001"
    assert result.preview.displayed[0].identity_status == "기존 소하천"


def test_component_code_maps_for_new_stream(tmp_path):
    db = tmp_path / "research.sqlite3"
    workbook = tmp_path / "구성코드.xlsx"
    workspace = tmp_path / "workspace"
    initialize_database(db)
    _bootstrap(db)
    _identity_workbook(
        workbook,
        ["시도코드", "시군구코드", "읍면동코드", "일련번호", "소하천명", "유역면적"],
        ["01", "234", "567", "001", "신규 합성 하천", 2.5],
    )
    source = _source(workbook, db, workspace)
    workflow = ImportMappingWorkflowService(db, workspace)
    state = workflow.set_scope(workflow.start(source, 1), "NATIONAL_2024", 1)
    result = workflow.preview(state, 1)
    assert result.preview.ready_for_import_preparation
    assert result.preview.displayed[0].stream_code == "01234567001"
    assert result.preview.displayed[0].identity_status == "신규 소하천"


def test_direct_component_mismatch_is_blocked_in_production_workflow(tmp_path):
    db = tmp_path / "research.sqlite3"
    workbook = tmp_path / "불일치.xlsx"
    workspace = tmp_path / "workspace"
    initialize_database(db)
    _bootstrap(db)
    _identity_workbook(
        workbook,
        [*(label for _name, label in CORE), "유역면적"],
        ["01234567002", "01", "234", "567", "001", "불일치 하천", 2.5],
    )
    source = _source(workbook, db, workspace)
    workflow = ImportMappingWorkflowService(db, workspace)
    state = workflow.set_scope(workflow.start(source, 1), "NATIONAL_2024", 1)
    result = workflow.preview(state, 1)
    assert not result.preview.ready_for_import_preparation
    assert result.preview.blocked == 1
    assert any("SOURCE_COMPONENT_MISMATCH" in issue for issue in result.preview.displayed[0].issues)


def test_existing_stream_name_mismatch_does_not_rename_or_change_identity(tmp_path):
    db = tmp_path / "research.sqlite3"
    workbook = tmp_path / "이름불일치.xlsx"
    workspace = tmp_path / "workspace"
    initialize_database(db)
    _bootstrap(db)
    with closing(connect_database(db)) as connection:
        connection.execute(
            "INSERT INTO small_stream (stream_code,province_code,city_county_code,town_code,"
            "stream_serial_no,stream_name,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?)",
            (
                "01234567001",
                "01",
                "234",
                "567",
                "001",
                "DB 원본명",
                "2026-09-23T00:00:00Z",
                "2026-09-23T00:00:00Z",
            ),
        )
    _identity_workbook(
        workbook,
        ["관리코드", "소하천명", "유역면적"],
        ["01234567001", "Excel 다른 이름", 2.5],
    )
    source = _source(workbook, db, workspace)
    before = _snapshot(db)
    workflow = ImportMappingWorkflowService(db, workspace)
    state = workflow.set_scope(workflow.start(source, 1), "NATIONAL_2024", 1)
    result = workflow.preview(state, 1)
    assert result.preview.ready_for_import_preparation
    assert result.preview.displayed[0].identity_status == "기존 소하천"
    assert _snapshot(db) == before


def test_old_unmapped_workspace_requires_explicit_auto_mapping(tmp_path, monkeypatch):
    db = tmp_path / "research.sqlite3"
    workbook = tmp_path / "기존작업.xlsx"
    workspace = tmp_path / "workspace"
    initialize_database(db)
    original = ImportCoreDictionaryBootstrapService.bootstrap
    with closing(connect_database(db)) as connection:
        monkeypatch.setattr(ImportCoreDictionaryBootstrapService, "bootstrap", lambda self: None)
        ResearchDictionaryBootstrapService(connection).bootstrap()
    _workbook(workbook)
    source = _source(workbook, db, workspace)
    workflow = ImportMappingWorkflowService(db, workspace)
    old = workflow.set_scope(workflow.start(source, 1), "NATIONAL_2024", 1)
    assert all(mapping.mapping_status == MappingStatus.UNMAPPED for mapping in old.mappings[:6])

    monkeypatch.setattr(ImportCoreDictionaryBootstrapService, "bootstrap", original)
    with closing(connect_database(db)) as connection:
        ImportCoreDictionaryBootstrapService(connection).bootstrap()
    resumed = workflow.start(old.source, 1)
    assert all(mapping.mapping_status == MappingStatus.UNMAPPED for mapping in resumed.mappings[:6])
    refreshed = workflow.update(resumed, 0, "auto", None, 1)
    assert all(
        mapping.mapping_status == MappingStatus.AUTO_MAPPED for mapping in refreshed.mappings[:7]
    )


def test_preview_is_bounded_safe_and_read_only(tmp_path):
    db = tmp_path / "research.sqlite3"
    workbook = tmp_path / "합성.xlsx"
    workspace = tmp_path / "workspace"
    initialize_database(db)
    _bootstrap(db)
    _workbook(workbook, invalid=True)
    source = _source(workbook, db, workspace)
    workflow = ImportMappingWorkflowService(db, workspace)
    before = _snapshot(db)
    state = workflow.set_scope(workflow.start(source, 1), "NATIONAL_2024", 1)
    result = workflow.preview(state, 1)
    assert result.preview.total == 2
    assert len(result.preview.displayed) == 2
    assert result.preview.displayed[0].stream_code == "01234567001"
    assert result.preview.blocked >= 1
    assert not result.preview.ready_for_import_preparation
    assert "가상값" not in repr(result.preview)
    assert _snapshot(db) == before
    changed = workflow.update(result, 9, "exclude", None, 1)
    assert changed.preview is None


def test_large_preview_projects_only_first_50_rows(tmp_path):
    db = tmp_path / "research.sqlite3"
    workbook = tmp_path / "합성.xlsx"
    workspace = tmp_path / "workspace"
    initialize_database(db)
    _bootstrap(db)
    _workbook(workbook)
    book = Workbook()
    sheet = book.active
    sheet.title = "합성 자료"
    sheet.append([*(label for _name, label in CORE), "유역면적", "Service Key", "미확정 열"])
    for number in range(65):
        code = f"01234567{number:03d}"
        sheet.append([code, "01", "234", "567", f"{number:03d}", "합성 하천", 2.5, "가상값", "A"])
    book.save(workbook)
    book.close()
    source = _source(workbook, db, workspace)
    workflow = ImportMappingWorkflowService(db, workspace)
    before = _snapshot(db)
    state = workflow.set_scope(workflow.start(source, 1), "NATIONAL_2024", 1)
    result = workflow.preview(state, 1)
    assert result.preview.total == 65
    assert len(result.preview.displayed) == 50
    assert result.preview.ready == 65
    assert "가상값" not in repr(result.preview)
    assert _snapshot(db) == before


def test_changed_source_and_inconsistent_dictionary_are_blocked(tmp_path):
    db = tmp_path / "research.sqlite3"
    workbook = tmp_path / "합성.xlsx"
    workspace = tmp_path / "workspace"
    initialize_database(db)
    _bootstrap(db)
    _workbook(workbook)
    source = _source(workbook, db, workspace)
    workflow = ImportMappingWorkflowService(db, workspace)
    state = workflow.start(source, 1)
    with closing(connect_database(db)) as connection:
        repository = DictionaryRepository(connection)
        item = repository.get_item_by_internal_name("basin_area")
        DictionaryService(repository).deactivate_item(item.dictionary_id)
    with pytest.raises(MappingWorkflowError, match="연구 사전"):
        workflow.start(state.source, 1)
    _workbook(workbook, invalid=True)
    with pytest.raises(MappingWorkflowError):
        workflow.start(state.source, 1)


def test_gui_10b_to_10c_preview_back_and_logout(tmp_path):
    app = QApplication.instance() or QApplication([])
    db = tmp_path / "gui.sqlite3"
    workspace = tmp_path / "workspace"
    workbook = tmp_path / "합성.xlsx"
    _workbook(workbook)
    controller = ApplicationController(db, workspace)
    try:
        _bootstrap(db)
        controller.auth_service.create_user(
            "synthetic_actor", "synthetic-password", "합성 작업자", "합성 부서", "연구자"
        )
        controller.start()
        controller.login_window.login_id.setText("synthetic_actor")
        controller.login_window.password.setText("synthetic-password")
        controller.login_window.submit.click()
        _wait(app, lambda: controller.main_window is not None)
        window = controller.main_window
        before = _snapshot(db)
        window.navigate("Excel 가져오기")
        structure = window.import_workspace
        assert not structure.next_button.isEnabled()
        structure.select_file(str(workbook))
        _wait(app, lambda: structure.state is not None)
        structure.sheets.selectRow(0)
        structure.structure_button.click()
        _wait(app, lambda: bool(structure.state and structure.state.columns))
        assert not structure.next_button.isEnabled()
        structure.save_button.click()
        _wait(app, lambda: structure.next_button.isEnabled())
        structure.next_button.click()
        mapping = window.import_mapping
        assert window.stack.currentWidget() is mapping
        _wait(app, lambda: mapping.state is not None)
        mapping.scope.setCurrentIndex(1)
        mapping.scope_button.click()
        _wait(
            app, lambda: mapping.state is not None and mapping.state.source_scope == "NATIONAL_2024"
        )
        assert mapping.mapping_table.rowCount() == 9
        assert mapping.state.rows[7].header == "제외 대상 열"
        mapping.mapping_table.selectRow(8)
        mapping.exclude_button.click()
        _wait(app, lambda: mapping.state.rows[8].status == "가져오지 않음")
        mapping.preview_button.click()
        _wait(app, lambda: mapping.state.preview is not None)
        assert mapping.state.preview.total == 1
        assert mapping.preview_table.rowCount() == 1
        assert mapping.state.preview.ready_for_import_preparation
        assert mapping.next_button.isEnabled()
        window.navigate("홈")
        window.navigate("Excel 가져오기")
        assert window.stack.currentWidget() is mapping
        _wait(app, lambda: mapping.state is not None)
        assert mapping.state.mappings[8].mapping_status == MappingStatus.DO_NOT_MAP
        assert mapping.state.preview is None
        for width, height in ((1440, 900), (1200, 800), (900, 600)):
            window.resize(width, height)
            app.processEvents()
            assert mapping.scroll_area.horizontalScrollBar().maximum() == 0
            if width == 900:
                assert mapping.mapping_table.horizontalScrollBar().maximum() > 0
                assert mapping.preview_table.horizontalScrollBar().maximum() > 0
        mapping.back_button.click()
        assert window.stack.currentWidget() is structure
        assert structure.next_button.isEnabled()
        assert _snapshot(db) == before
        window.logout_button.click()
        app.processEvents()
        assert controller.session is None
    finally:
        controller.close()
        app.processEvents()


def test_preview_stale_result_cannot_restore_old_mapping(tmp_path):
    app = QApplication.instance() or QApplication([])
    db = tmp_path / "research.sqlite3"
    workspace = tmp_path / "workspace"
    workbook = tmp_path / "합성.xlsx"
    initialize_database(db)
    _bootstrap(db)
    _workbook(workbook)
    source = _source(workbook, db, workspace)
    initial = ImportMappingWorkflowService(db, workspace).start(source, 1)

    class PendingPool:
        def __init__(self):
            self.tasks = []

        def start(self, task):
            self.tasks.append(task)

    view = ImportMappingView(db, 1, workspace)
    view.show()
    view.state = initial
    view._show_mapping()
    view._controls()
    pool = PendingPool()
    view.pool = pool
    view._preview()
    view.mapping_table.selectRow(8)
    view._change("exclude")
    pool.tasks[1].run()
    app.processEvents()
    assert view.state.mappings[8].mapping_status == MappingStatus.DO_NOT_MAP
    pool.tasks[0].run()
    app.processEvents()
    assert view.state.preview is None
    view._preview()
    view.deactivate()
    pool.tasks[2].run()
    app.processEvents()
    assert view.state.preview is None
    view.close()
