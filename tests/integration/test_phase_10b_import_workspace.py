"""합성 Excel과 임시 DB/Workspace로 10B GUI 진입·재개 경계를 검증한다."""

import os
import time
from contextlib import closing

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication, QLabel

from small_stream_research_tool.app.controller import ApplicationController
from small_stream_research_tool.database import connect_database, initialize_database
from small_stream_research_tool.database.connection import transaction
from small_stream_research_tool.models.import_workflow import ImportInspectionError
from small_stream_research_tool.services import import_inspection_service
from small_stream_research_tool.services.import_inspection_service import (
    ImportInspectionService,
)
from small_stream_research_tool.utils.file_hash import file_sha256

STAMP = "2026-09-22T00:00:00Z"


@pytest.fixture(scope="session")
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "합성 소하천.xlsx"
    book = Workbook()
    main = book.active
    main.title = "연구 특성"
    main.merge_cells("A1:B1")
    main["A1"] = "합성 그룹"
    main["A2"] = "관리코드"
    main["B2"] = "유역면적"
    main["A3"] = "01234567001"
    main["B3"] = 12.5
    hidden = book.create_sheet("숨김 시트")
    hidden.sheet_state = "hidden"
    hidden.append(["합성 헤더"])
    hidden.append(["합성 값"])
    book.create_sheet("빈 시트")
    chart = BarChart()
    chart.add_data(Reference(main, min_col=2, min_row=2, max_row=3), titles_from_data=True)
    book.create_chartsheet("차트").add_chart(chart)
    book.save(path)
    book.close()
    return path


def wait_for(app, predicate, seconds=10):
    end = time.monotonic() + seconds
    while time.monotonic() < end:
        app.processEvents()
        if predicate():
            return
        time.sleep(0.02)
    pytest.fail("합성 workbook 작업이 끝나지 않았습니다.")


def snapshot(path):
    with closing(connect_database(path)) as connection:
        return tuple(connection.iterdump())


def test_inspect_structure_workspace_resume_and_db_read_only(tmp_path, workbook):
    db = tmp_path / "research.sqlite3"
    initialize_database(db)
    service = ImportInspectionService(db, tmp_path / "workspace")
    before = snapshot(db)
    inspected = service.inspect(workbook)
    assert inspected.file_hash == file_sha256(workbook)
    assert inspected.file_name == workbook.name
    assert inspected.file_size > 0 and len(inspected.sheets) == 4
    assert [sheet.name for sheet in inspected.sheets] == [
        "연구 특성",
        "숨김 시트",
        "빈 시트",
        "차트",
    ]
    assert inspected.sheets[1].selectable and inspected.sheets[1].state == "hidden"
    assert not inspected.sheets[2].selectable
    assert not inspected.sheets[3].selectable and inspected.sheets[3].kind == "chartsheet"
    with pytest.raises(ImportInspectionError):
        service.structure(inspected, "차트", 1, 1, 2)
    with pytest.raises(ImportInspectionError):
        service.structure(inspected, "연구 특성", 2, 1, 3)
    with pytest.raises(ImportInspectionError):
        service.structure(inspected, "연구 특성", 1, 2, 2)
    structured = service.structure(inspected, "연구 특성", 1, 2, 3)
    assert len(structured.columns) == 2
    assert structured.columns[0].text == "합성 그룹 | 관리코드"
    assert structured.columns[1].text == "합성 그룹 | 유역면적"
    assert "01234567001" not in repr(structured.columns)
    saved = service.save(structured, 1)
    assert saved.workspace_saved_at
    resumed = service.resume(1)
    assert resumed.selected_sheet == "연구 특성"
    assert (resumed.header_start_row, resumed.header_end_row, resumed.data_start_row) == (1, 2, 3)
    assert resumed.columns == structured.columns
    assert snapshot(db) == before
    assert "01234567001" not in (tmp_path / "workspace" / "user_1.json").read_text(encoding="utf-8")
    book = Workbook()
    book.active.append(["변경된 합성 파일"])
    book.save(workbook)
    book.close()
    with pytest.raises(ImportInspectionError, match="원본 파일"):
        service.resume(1)


def test_duplicate_hash_safe_summary_and_extension(tmp_path, workbook):
    db = tmp_path / "research.sqlite3"
    initialize_database(db)
    digest = file_sha256(workbook)
    with closing(connect_database(db)) as connection:
        with transaction(connection):
            source = connection.execute(
                "INSERT INTO source_file (file_name,original_path,file_hash,registered_at) "
                "VALUES (?,?,?,?)",
                ("C:/PRIVATE/과거 파일.xlsx", "C:/PRIVATE/원본.xlsx", digest, STAMP),
            ).lastrowid
            connection.execute(
                "INSERT INTO import_history "
                "(source_file_id,batch_code,import_type,status,started_at,created_at) "
                "VALUES (?,?,?,?,?,?)",
                (source, "synthetic-previous", "EXCEL", "SUCCESS", STAMP, STAMP),
            )
    before = snapshot(db)
    state = ImportInspectionService(db, tmp_path / "workspace").inspect(workbook)
    assert len(state.duplicates) == 1
    assert state.duplicates[0].file_name == "과거 파일.xlsx"
    assert state.duplicates[0].status == "SUCCESS"
    assert "C:/PRIVATE" not in repr(state)
    assert snapshot(db) == before
    unsupported = tmp_path / "synthetic.csv"
    unsupported.write_text("합성", encoding="utf-8")
    with pytest.raises(ImportInspectionError, match="xlsx"):
        ImportInspectionService(db, tmp_path / "workspace").inspect(unsupported)


def test_large_workbook_inspects_only_header_columns(tmp_path):
    db = tmp_path / "research.sqlite3"
    initialize_database(db)
    path = tmp_path / "대량 합성.xlsx"
    book = Workbook(write_only=True)
    sheet = book.create_sheet("대량")
    sheet.append(["관리코드", "값"])
    for number in range(3000):
        sheet.append([f"01234567{number:03d}", number])
    book.save(path)
    book.close()
    service = ImportInspectionService(db, tmp_path / "workspace")
    state = service.inspect(path)
    assert state.sheets[0].row_count == 3001
    structured = service.structure(state, "대량", 1, 1, 2)
    assert len(structured.columns) == 2
    assert not hasattr(structured, "rows")


def test_hash_metadata_change_and_unreadable_workbook_are_safe(tmp_path, workbook, monkeypatch):
    db = tmp_path / "research.sqlite3"
    initialize_database(db)
    service = ImportInspectionService(db, tmp_path / "workspace")
    original_hash = import_inspection_service.file_sha256

    def change_during_hash(path):
        digest = original_hash(path)
        with open(path, "ab") as target:
            target.write(b"synthetic-change")
        return digest

    monkeypatch.setattr(import_inspection_service, "file_sha256", change_during_hash)
    with pytest.raises(ImportInspectionError, match="변경") as error:
        service.inspect(workbook)
    assert str(workbook) not in str(error.value)
    monkeypatch.setattr(import_inspection_service, "file_sha256", original_hash)
    damaged = tmp_path / "손상 합성.xlsx"
    damaged.write_bytes(b"not-an-xlsx")
    with pytest.raises(ImportInspectionError) as error:
        service.inspect(damaged)
    assert str(damaged) not in str(error.value)


def test_missing_file_and_formatted_empty_sheet_are_not_selectable(tmp_path, workbook):
    db = tmp_path / "research.sqlite3"
    initialize_database(db)
    service = ImportInspectionService(db, tmp_path / "workspace")
    with pytest.raises(ImportInspectionError) as error:
        service.inspect(tmp_path / "없는 파일.xlsx")
    assert str(tmp_path) not in str(error.value)
    with pytest.raises(ImportInspectionError):
        service.inspect(tmp_path)

    book = Workbook()
    sheet = book.active
    sheet.title = "서식만 있는 시트"
    sheet["C5"].number_format = "0.00"
    filled = book.create_sheet("첫 행이 빈 시트")
    filled["A3"] = "합성 헤더"
    book.save(workbook)
    book.close()
    state = service.inspect(workbook)
    assert not state.sheets[0].selectable
    assert state.sheets[1].selectable


@pytest.fixture
def controller(tmp_path, app):
    control = ApplicationController(tmp_path / "gui.sqlite3", tmp_path / "workspace")
    control.auth_service.create_user(
        "synthetic_actor", "synthetic-password", "합성 작업자", "합성 부서", "연구자"
    )
    yield control
    control.close()
    app.processEvents()


def login(control, app):
    control.start()
    window = control.login_window
    window.login_id.setText("synthetic_actor")
    window.password.setText("synthetic-password")
    window.submit.click()
    wait_for(app, lambda: control.main_window is not None)
    return control.main_window


def visible_text(view):
    return " ".join(label.text() for label in view.findChildren(QLabel))


def test_gui_offscreen_workflow_navigation_resume_logout(controller, workbook, app):
    window = login(controller, app)
    before = snapshot(controller.db_path)
    window.nav_buttons["Excel 가져오기"].click()
    view = window.import_workspace
    assert window.stack.currentWidget() is view
    assert window.nav_buttons["Excel 가져오기"].objectName() == "navSelected"
    assert not view.next_button.isEnabled()
    view.select_file(str(workbook))
    assert "확인하는 중" in view.status.text()
    wait_for(app, lambda: view.state is not None)
    assert view.file_name.text() == workbook.name
    assert "SHA-256 확인 완료" in view.file_info.text()
    assert view.state.file_hash not in visible_text(view)
    assert not view.duplicate_warning.isVisible()
    assert str(workbook) not in visible_text(view)
    view.sheets.selectRow(0)
    view.header_start.setValue(1)
    view.header_end.setValue(2)
    view.data_start.setValue(3)
    view.structure_button.click()
    wait_for(app, lambda: bool(view.state and view.state.columns))
    assert view.headers.rowCount() == 2
    assert "01234567001" not in visible_text(view)
    for width, height in ((1440, 900), (1200, 800), (900, 600)):
        window.resize(width, height)
        app.processEvents()
        assert view.scroll_area.viewport().width() > 0
        assert view.scroll_area.horizontalScrollBar().maximum() == 0
        for control in (
            view.choose_button,
            view.resume_button,
            view.sheets,
            view.header_start,
            view.data_start,
            view.headers,
            view.save_button,
        ):
            view.scroll_area.ensureWidgetVisible(control)
            app.processEvents()
            assert control.isVisible()
    assert view.headers.horizontalScrollBar().maximum() > 0
    view.save_button.click()
    wait_for(app, lambda: bool(view.state and view.state.workspace_saved_at))
    assert not view.next_button.isEnabled()
    window.navigate("홈")
    assert window.stack.currentWidget() is window.home
    window.navigate("소하천 조회")
    window.navigate("Excel 가져오기")
    assert window.stack.currentWidget() is view
    assert view.state is not None
    view.resume_button.click()
    wait_for(app, lambda: view.status.text() == "저장된 작업을 재개했습니다.")
    assert view.state.columns and view.state.workspace_saved_at
    assert snapshot(controller.db_path) == before
    window.logout_button.click()
    app.processEvents()
    assert controller.session is None and controller.login_window.isVisible()


def test_gui_duplicate_warning_and_sheet_selection_policy(controller, workbook, app):
    window = login(controller, app)
    digest = file_sha256(workbook)
    with closing(connect_database(controller.db_path)) as connection:
        with transaction(connection):
            source = connection.execute(
                "INSERT INTO source_file (file_name,original_path,file_hash,registered_at) "
                "VALUES (?,?,?,?)",
                ("C:/PRIVATE/다른 이름.xlsx", "C:/PRIVATE/원본.xlsx", digest, STAMP),
            ).lastrowid
            connection.execute(
                "INSERT INTO import_history "
                "(source_file_id,batch_code,import_type,status,started_at,created_at) "
                "VALUES (?,?,?,?,?,?)",
                (source, "prior-synthetic", "EXCEL", "SUCCESS", STAMP, STAMP),
            )
    before = snapshot(controller.db_path)
    window.navigate("Excel 가져오기")
    view = window.import_workspace
    view.select_file(str(workbook))
    wait_for(app, lambda: view.state is not None)
    assert "동일한 내용" in view.duplicate_warning.text()
    assert "다른 이름.xlsx" in view.duplicate_warning.text()
    assert "C:/PRIVATE" not in visible_text(view)
    assert not (view.sheets.item(2, 0).flags() & Qt.ItemFlag.ItemIsSelectable)
    assert not (view.sheets.item(3, 0).flags() & Qt.ItemFlag.ItemIsSelectable)
    view.sheets.selectRow(1)
    assert view.structure_button.isEnabled()
    view.header_start.setValue(1)
    view.header_end.setValue(1)
    view.data_start.setValue(2)
    view.structure_button.click()
    wait_for(app, lambda: bool(view.state and view.state.columns))
    assert view.state.selected_sheet == "숨김 시트"
    assert view.headers.rowCount() == 1
    assert snapshot(controller.db_path) == before


def test_gui_changed_source_resume_is_safe(controller, workbook, app):
    window = login(controller, app)
    window.navigate("Excel 가져오기")
    view = window.import_workspace
    view.select_file(str(workbook))
    wait_for(app, lambda: view.state is not None)
    view.sheets.selectRow(0)
    view.header_end.setValue(2)
    view.data_start.setValue(3)
    view.structure_button.click()
    wait_for(app, lambda: bool(view.state and view.state.columns))
    view.save_button.click()
    wait_for(app, lambda: bool(view.state and view.state.workspace_saved_at))
    book = Workbook()
    book.active.append(["변경된 합성 파일"])
    book.save(workbook)
    book.close()
    view.resume_button.click()
    wait_for(app, lambda: "재개할 수 없습니다" in view.status.text())
    assert view.state is None
    assert not view.save_button.isEnabled()
    assert view.sheets.rowCount() == 0
    assert str(workbook) not in visible_text(view)
    assert "Traceback" not in visible_text(view)


def test_gui_missing_source_resume_preserves_workspace(controller, workbook, app):
    window = login(controller, app)
    window.navigate("Excel 가져오기")
    view = window.import_workspace
    view.select_file(str(workbook))
    wait_for(app, lambda: view.state is not None)
    view.sheets.selectRow(0)
    view.header_end.setValue(2)
    view.data_start.setValue(3)
    view.structure_button.click()
    wait_for(app, lambda: bool(view.state and view.state.columns))
    view.save_button.click()
    wait_for(app, lambda: bool(view.state and view.state.workspace_saved_at))
    workspace_path = controller.workspace_dir / f"user_{controller.session.user_id}.json"
    assert workspace_path.is_file()
    workbook.unlink()
    view.resume_button.click()
    wait_for(app, lambda: "재개할 수 없습니다" in view.status.text())
    assert view.state is None
    assert not view.save_button.isEnabled()
    assert workspace_path.is_file()
    assert str(workbook) not in visible_text(view)


def test_gui_stale_file_and_navigation_results(controller, workbook, tmp_path, app):
    window = login(controller, app)
    view = window.import_workspace
    window.navigate("Excel 가져오기")
    second = tmp_path / "다른 합성.xlsx"
    book = Workbook()
    book.active.append(["다른 헤더"])
    book.save(second)
    book.close()

    class PendingPool:
        def __init__(self):
            self.tasks = []

        def start(self, task):
            self.tasks.append(task)

    pool = PendingPool()
    view.pool = pool
    view.select_file(str(workbook))
    view.select_file(str(second))
    assert len(pool.tasks) == 2
    pool.tasks[1].run()
    app.processEvents()
    assert view.file_name.text() == second.name
    pool.tasks[0].run()
    app.processEvents()
    assert view.file_name.text() == second.name
    view.select_file(str(workbook))
    window.navigate("홈")
    pool.tasks[2].run()
    app.processEvents()
    assert window.stack.currentWidget() is window.home
    assert view.file_name.text() == "파일 확인 중입니다."
    assert view.state is None


def test_gui_safe_error_and_no_local_path(controller, workbook, app):
    window = login(controller, app)
    window.navigate("Excel 가져오기")
    view = window.import_workspace
    unsupported = workbook.with_suffix(".xlsm")
    unsupported.write_text("합성", encoding="utf-8")
    view.select_file(str(unsupported))
    wait_for(app, lambda: ".xlsx" in view.status.text())
    assert str(unsupported) not in visible_text(view)
    assert "Traceback" not in visible_text(view)
    assert not view.save_button.isEnabled()
