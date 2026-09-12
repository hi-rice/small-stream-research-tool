"""Synthetic Preview와 임시 SQLite만 사용한다. 운영 DB와 실제 연구자료는 사용하지 않는다."""

import json
import sqlite3
from contextlib import closing
from dataclasses import asdict, replace
from datetime import date, datetime
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from small_stream_research_tool.database import connect_database, initialize_database
from small_stream_research_tool.models.column_mapping import (
    ColumnMappingDraft,
    MappingHeaderPart,
    MappingMethod,
    MappingStatus,
)
from small_stream_research_tool.models.excel import ExcelCell, ExcelRow
from small_stream_research_tool.models.import_preparation import (
    ImportFieldPolicy,
    PreparationStatus,
    RowAction,
)
from small_stream_research_tool.models.import_preparation_errors import (
    InvalidPreparationArgumentError,
)
from small_stream_research_tool.models.import_preview import PreviewFieldPolicy, PreviewStatus
from small_stream_research_tool.models.workspace import WorkspaceStep
from small_stream_research_tool.repositories.dictionary_repository import DictionaryRepository
from small_stream_research_tool.repositories.stream_lookup_repository import StreamLookupRepository
from small_stream_research_tool.services.dictionary_service import DictionaryService
from small_stream_research_tool.services.import_preparation_service import ImportPreparationService
from small_stream_research_tool.services.import_preview_service import ImportPreviewService
from small_stream_research_tool.services.workspace_service import WorkspaceService
from small_stream_research_tool.utils.file_hash import file_sha256

IDENTITY = (
    "stream_code",
    "province_code",
    "city_county_code",
    "town_code",
    "stream_serial_no",
    "stream_name",
)
OPTIONAL_TEXT = (
    "province_name",
    "city_county_name",
    "town_name",
    "river_system",
    "source_address",
    "end_address",
)
COORDINATES = ("source_latitude", "source_longitude", "end_latitude", "end_longitude")
DEFAULT_NAMES = (*IDENTITY, "synthetic_real")
DEFAULT_VALUES = ("12345678008", "12", "345", "678", "008", "Synthetic Stream A", " 4.24 ")


@pytest.fixture
def context(tmp_path):
    path = tmp_path / "synthetic.sqlite3"
    initialize_database(path)
    with closing(connect_database(path)) as connection:
        dictionary = DictionaryService(DictionaryRepository(connection))
        category = dictionary.create_category("synthetic", "Synthetic Category")
        version = dictionary.create_version("synthetic-v1")
        unit = dictionary.create_unit("Synthetic Unit A", "synA")
        other = dictionary.create_unit("Synthetic Unit B", "synB")
        dictionary.register_conversion(
            unit.unit_id, other.unit_id, factor=10, offset=1, approved=True
        )
        definitions = dict.fromkeys((*IDENTITY, *OPTIONAL_TEXT, "synthetic_hidden"), "TEXT")
        definitions |= dict.fromkeys((*COORDINATES, "synthetic_real"), "REAL")
        definitions |= {
            "synthetic_integer": "INTEGER",
            "synthetic_date": "DATE",
            "synthetic_datetime": "DATETIME",
            "created_at": "TEXT",
        }
        items = {}
        for name, dtype in definitions.items():
            items[name] = dictionary.create_item(
                "Synthetic " + name,
                name,
                category.category_id,
                dtype,
                version.version_id,
                unit_id=unit.unit_id if name == "synthetic_real" else None,
                required=True,
                nullable=False,
            )
        connection.execute(
            """INSERT INTO small_stream
            (stream_code, province_code, city_county_code, town_code, stream_serial_no,
             stream_name, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                "12345678009",
                "12",
                "345",
                "678",
                "009",
                "Synthetic Existing Name",
                "2020-01-01T00:00:00Z",
                "2020-01-01T00:00:00Z",
            ),
        )
        yield SimpleNamespace(
            connection=connection,
            dictionary=dictionary,
            items=items,
            unit=unit,
            preview=ImportPreviewService(dictionary, StreamLookupRepository(connection)),
            service=ImportPreparationService(dictionary),
        )


def input_row(values=DEFAULT_VALUES, source_row=2):
    return ExcelRow(
        source_row,
        tuple(
            ExcelCell(
                source_row,
                index,
                get_column_letter(index),
                value,
                "s" if isinstance(value, str) else "n",
                False,
                "General",
            )
            for index, value in enumerate(values, 1)
        ),
    )


def mappings(context, names=DEFAULT_NAMES, *, header="Synthetic Header"):
    return tuple(
        ColumnMappingDraft(
            index,
            get_column_letter(index),
            header,
            (MappingHeaderPart(1, header, None, None, None, None),),
            context.items[name].dictionary_id,
            MappingMethod.USER,
            MappingStatus.USER_MAPPED,
        )
        for index, name in enumerate(names, 1)
    )


def preview(context, values=DEFAULT_VALUES, names=DEFAULT_NAMES, **options):
    return context.preview.build_preview(
        [input_row(values)], mappings(context, names), **options
    ).rows[0]


def prepare(context, values=DEFAULT_VALUES, names=DEFAULT_NAMES, **options):
    return context.service.prepare_row(preview(context, values, names), **options)


def test_new_stream_candidate_and_provenance(context):
    result = prepare(context)
    assert result.status == PreparationStatus.READY and result.row_action == RowAction.CREATE_STREAM
    assert result.source_row == 2 and result.stream_code == "12345678008"
    assert result.core_data.stream_name == "Synthetic Stream A"
    assert (
        result.core_data.province_code,
        result.core_data.city_county_code,
        result.core_data.town_code,
        result.core_data.stream_serial_no,
    ) == ("12", "345", "678", "008")
    assert len(result.prepared_values) == 1
    value = result.prepared_values[0]
    assert value.dictionary_id == context.items["synthetic_real"].dictionary_id
    assert (value.source_row, value.source_column_index, value.value_number) == (2, 7, 4.24)
    assert value.original_value == " 4.24 " and value.original_unit is None
    assert value.unit_id == context.unit.unit_id  # 등록된 factor 10 / offset 1 미적용


@pytest.mark.parametrize("missing", [None, "", "  "])
def test_new_name_required_without_placeholder(context, missing):
    result = prepare(context, (*DEFAULT_VALUES[:5], missing, 4))
    assert result.status == PreparationStatus.BLOCKED and result.core_data is None
    assert result.stream_name is None
    assert "STREAM_NAME_REQUIRED" in {i.code for i in result.issues}


@pytest.mark.parametrize("name", [None, "", "   ", "Synthetic Different Name"])
def test_existing_preserves_core_even_without_name(context, name):
    result = prepare(context, ("12345678009", "12", "345", "678", "009", name, 4))
    assert result.status == PreparationStatus.READY
    assert result.row_action == RowAction.USE_EXISTING_STREAM and result.core_data is None
    assert len(result.prepared_values) == 1
    assert (
        context.connection.execute("SELECT stream_name FROM small_stream").fetchone()[0]
        == "Synthetic Existing Name"
    )


def test_source_only_existing_and_new_components_required(context):
    existing = prepare(context, ("12345678009", 4), ("stream_code", "synthetic_real"))
    assert existing.status == PreparationStatus.READY
    new = prepare(
        context,
        ("12345678008", "Synthetic Stream A", 4),
        ("stream_code", "stream_name", "synthetic_real"),
    )
    assert new.status == PreparationStatus.BLOCKED
    assert "COMPONENTS_REQUIRED" in {i.code for i in new.issues}


def test_generated_code_and_confirmed_zero_format(context):
    row = input_row((None, "12", "345", "678", 8, "Synthetic Stream A", 4))
    cells = list(row.cells)
    cells[4] = replace(cells[4], number_format="000")
    before = replace(row, cells=tuple(cells))
    result = context.service.prepare_row(
        context.preview.build_preview([before], mappings(context)).rows[0]
    )
    assert result.status == PreparationStatus.READY
    assert result.core_data.stream_serial_no == "008" and result.stream_code == "12345678008"
    assert before.cells[4].value == 8


@pytest.mark.parametrize("value", [None, "", "  "])
def test_missing_characteristic_not_failure(context, value):
    result = prepare(context, (*DEFAULT_VALUES[:6], value))
    assert result.status == PreparationStatus.READY and result.prepared_values == ()
    assert not any(i.blocking for i in result.issues)
    assert "CONVERSION_FAILED" not in {i.code for i in result.issues}
    assert "MISSING_VALUE" in {i.code for i in result.issues}


def test_existing_no_op(context):
    result = prepare(context, ("12345678009",), ("stream_code",))
    assert result.status == PreparationStatus.BLOCKED
    assert result.row_action is None and result.prepared_values == ()
    assert "NO_CHARACTERISTIC_VALUES" in {i.code for i in result.issues}


@pytest.mark.parametrize("new", [True, False])
def test_invalid_mapped_value_blocks_row(context, new, caplog):
    values = list(DEFAULT_VALUES)
    if not new:
        values[0], values[4] = "12345678009", "009"
    values[6] = "SYNTHETIC_INVALID_RAW_4.24 km"
    result = prepare(context, values)
    assert result.status == PreparationStatus.BLOCKED and result.prepared_values == ()
    issue = next(i for i in result.issues if i.code == "CONVERSION_FAILED")
    assert issue.blocking and (issue.source_row, issue.source_column_index) == (2, 7)
    assert values[6] not in repr(result) + caplog.text


@pytest.mark.parametrize("kind,code", [("f", "FORMULA_CELL"), ("e", "EXCEL_ERROR_CELL")])
def test_formula_error_mapped_characteristic(context, kind, code):
    row = input_row()
    row = replace(
        row, cells=(*row.cells[:6], replace(row.cells[6], value="SYNTHETIC_RAW", value_type=kind))
    )
    result = context.service.prepare_row(
        context.preview.build_preview([row], mappings(context)).rows[0]
    )
    assert result.status == PreparationStatus.BLOCKED and result.prepared_values == ()
    assert code in {i.code for i in result.issues}


def test_preview_blocking_never_released(context):
    before = preview(context, ("invalid", *DEFAULT_VALUES[1:]))
    assert before.status == PreviewStatus.NEEDS_REVIEW
    result = context.service.prepare_row(before)
    assert result.status == PreparationStatus.BLOCKED and not result.prepared_values


def test_excluded_has_priority_and_no_values(context):
    before = preview(context, ("invalid", *DEFAULT_VALUES[1:]), excluded_rows={2})
    result = context.service.prepare_row(before)
    assert result.status == PreparationStatus.EXCLUDED
    assert result.core_data is None and result.prepared_values == () and result.stream_code is None


def test_optional_core_fields_separated_and_range_checked(context):
    names = (*DEFAULT_NAMES, *OPTIONAL_TEXT, *COORDINATES)
    values = (*DEFAULT_VALUES, *("Synthetic Core" for _ in OPTIONAL_TEXT), 45, 120, -45, -120)
    result = prepare(context, values, names)
    assert result.status == PreparationStatus.READY and len(result.prepared_values) == 1
    assert len(result.core_data.optional_fields) == 10
    assert dict(result.core_data.optional_fields)["source_latitude"] == 45.0
    failed = prepare(context, (*values[:-4], 91, 120, -45, -120), names)
    assert failed.status == PreparationStatus.BLOCKED and failed.core_data is None
    assert "CORE_COORDINATE_RANGE" in {i.code for i in failed.issues}


def test_existing_optional_core_not_update_or_characteristic(context):
    result = prepare(
        context,
        ("12345678009", 4, "Synthetic Replacement", 999),
        ("stream_code", "synthetic_real", "province_name", "source_latitude"),
    )
    assert result.status == PreparationStatus.READY and result.core_data is None
    assert len(result.prepared_values) == 1


def test_import_policy_separate_from_display_policy(context):
    names = (*DEFAULT_NAMES, "synthetic_hidden")
    values = (*DEFAULT_VALUES, "SYNTHETIC_PRIVATE_VALUE")
    before = preview(
        context,
        values,
        names,
        field_policy=PreviewFieldPolicy(
            masked_internal_names=frozenset({"synthetic_real", "stream_name"})
        ),
    )
    assert before.stream_name == "[MASKED]"
    result = context.service.prepare_row(
        before, field_policy=ImportFieldPolicy(frozenset({"synthetic_hidden"}))
    )
    assert result.status == PreparationStatus.READY and len(result.prepared_values) == 1
    assert result.core_data.stream_name == "Synthetic Stream A"
    assert result.prepared_values[0].value_number == 4.24
    assert "SYNTHETIC_PRIVATE_VALUE" not in str(asdict(result))
    assert ImportFieldPolicy().excluded_internal_names == frozenset()


def test_policy_runs_before_conversion_and_uses_internal_name(context):
    before = preview(context, (*DEFAULT_VALUES[:6], "bad"))
    result = context.service.prepare_row(
        before, field_policy=ImportFieldPolicy(frozenset({"synthetic_real"}))
    )
    assert result.status == PreparationStatus.READY and result.prepared_values == ()
    assert "CONVERSION_FAILED" not in {i.code for i in result.issues}
    result = context.service.prepare_row(
        before, field_policy=ImportFieldPolicy(frozenset({"Synthetic Header"}))
    )
    assert result.status == PreparationStatus.BLOCKED


def test_policy_excluded_name_cannot_create_placeholder(context):
    result = prepare(context, field_policy=ImportFieldPolicy(frozenset({"stream_name"})))
    assert result.status == PreparationStatus.BLOCKED and result.stream_name is None


def test_header_does_not_supply_units_or_exclusion(context):
    before = context.preview.build_preview(
        [input_row()], mappings(context, header="Synthetic Unit [synB]")
    ).rows[0]
    result = context.service.prepare_row(before)
    assert result.prepared_values[0].original_unit is None
    assert result.prepared_values[0].unit_id == context.unit.unit_id


def test_dictionary_deactivated_after_preview(context):
    before = preview(context)
    context.connection.execute(
        "UPDATE data_dictionary SET is_active=0 WHERE dictionary_id=?",
        (context.items["synthetic_real"].dictionary_id,),
    )
    result = context.service.prepare_row(before)
    assert result.status == PreparationStatus.BLOCKED and not result.prepared_values
    assert "DICTIONARY_UNAVAILABLE" in {i.code for i in result.issues}


@pytest.mark.parametrize("kind", ["category", "unit", "semantic", "core_type"])
def test_dictionary_changes_fail_closed(context, kind):
    before = preview(context)
    if kind == "category":
        context.connection.execute("UPDATE data_category SET is_active=0")
    elif kind == "unit":
        context.connection.execute("UPDATE unit_dictionary SET is_active=0")
    elif kind == "semantic":
        context.connection.execute(
            "UPDATE data_dictionary SET internal_name='synthetic_changed' "
            "WHERE internal_name='synthetic_real'"
        )
    else:
        context.connection.execute(
            "UPDATE data_dictionary SET data_type='INTEGER' WHERE internal_name='stream_name'"
        )
    assert context.service.prepare_row(before).status == PreparationStatus.BLOCKED


def test_forged_duplicate_or_location_blocks(context):
    before = preview(context)
    duplicate = replace(before, mapped_values=(*before.mapped_values, before.mapped_values[-1]))
    assert context.service.prepare_row(duplicate).status == PreparationStatus.BLOCKED
    last = before.mapped_values[-1]
    wrong = replace(
        before, mapped_values=(*before.mapped_values[:-1], replace(last, source_column_index=9))
    )
    assert context.service.prepare_row(wrong).status == PreparationStatus.BLOCKED


def test_system_fields_not_characteristics(context):
    result = prepare(context, (*DEFAULT_VALUES, "Synthetic System"), (*DEFAULT_NAMES, "created_at"))
    assert result.status == PreparationStatus.BLOCKED
    assert len(result.prepared_values) == 1


def test_summary_counts_only_ready_actions_and_values(context):
    inputs = [
        input_row(),
        input_row(("12345678009", "12", "345", "678", "009", None, 3), 3),
        input_row((*DEFAULT_VALUES[:5], None, 5), 4),
        input_row(DEFAULT_VALUES, 5),
    ]
    before = context.preview.build_preview(inputs, mappings(context), excluded_rows={5})
    result = context.service.prepare(before.rows)
    summary = result.summary
    assert (
        summary.total_rows,
        summary.ready_rows,
        summary.blocked_rows,
        summary.excluded_rows,
    ) == (4, 2, 1, 1)
    assert (
        summary.create_stream_rows,
        summary.use_existing_stream_rows,
        summary.prepared_value_count,
    ) == (1, 1, 2)
    assert len(result.rows[2].prepared_values) == 1  # BLOCKED의 부분 후보는 summary에서 제외


def test_immutable_deterministic_and_no_db_writes(context):
    source, mapping = input_row(), mappings(context)
    before = context.preview.build_preview([source], mapping)
    snapshot = asdict(before)
    database_before = "\n".join(context.connection.iterdump())
    changes = context.connection.total_changes
    context.connection.execute("BEGIN")
    allowed = {sqlite3.SQLITE_SELECT, sqlite3.SQLITE_READ, sqlite3.SQLITE_FUNCTION}
    calls = []

    def authorizer(action, *_):
        calls.append(action)
        return sqlite3.SQLITE_OK if action in allowed else sqlite3.SQLITE_DENY

    context.connection.set_authorizer(authorizer)
    try:
        result = context.service.prepare(before.rows)
        assert result == context.service.prepare(before.rows)
        assert context.connection.in_transaction and context.connection.total_changes == changes
        assert set(calls) <= allowed
    finally:
        context.connection.set_authorizer(None)
        context.connection.rollback()
    assert "\n".join(context.connection.iterdump()) == database_before
    assert asdict(before) == snapshot and source == input_row() and mapping == mappings(context)
    for table in (
        "source_file",
        "import_history",
        "import_sheet",
        "import_column_mapping",
        "characteristic_value",
        "data_quality_issue",
        "record_history",
        "stream_characteristic",
    ):
        assert context.connection.execute(f"SELECT count(*) FROM {table}").fetchone()[0] == 0


def test_workspace_stays_metadata_only(context, tmp_path):
    path = tmp_path / "synthetic.xlsx"
    workbook = Workbook()
    workbook.active.append(list(DEFAULT_NAMES))
    workbook.active.append(list(DEFAULT_VALUES))
    workbook.save(path)
    workbook.close()
    original_hash = file_sha256(path)
    workspace = WorkspaceService(tmp_path / "workspace")
    draft = workspace.create_workspace(
        current_user_id=1,
        source_file_path=path,
        selected_sheet_name="Sheet",
        header_start_row=1,
        header_end_row=1,
        data_start_row=2,
        column_mappings=mappings(context),
        current_step=WorkspaceStep.PREVIEW,
    )
    workspace.save_workspace(draft, 1)
    saved_path = tmp_path / "workspace" / "user_1.json"
    saved = saved_path.read_bytes()
    before = context.preview.rebuild_from_workspace(workspace, 1)
    result = context.service.prepare(before.rows)
    assert result.summary.ready_rows == 1
    assert saved_path.read_bytes() == saved and file_sha256(path) == original_hash
    data = json.loads(saved)
    assert data["workspace_version"] == 1 and data["current_step"] == "PREVIEW"
    assert not {"rows", "prepared_values", "original_value", "core_data"} & data.keys()
    assert "Synthetic Stream A" not in saved.decode()


def test_duplicate_row_numbers_and_invalid_policy(context):
    before = preview(context)
    with pytest.raises(InvalidPreparationArgumentError):
        context.service.prepare([before, before])
    with pytest.raises(InvalidPreparationArgumentError):
        context.service.prepare_row(before, field_policy=PreviewFieldPolicy())


def test_all_typed_fields_and_missing_do_not_become_null_rows(context):
    names = (
        "stream_code",
        "synthetic_integer",
        "synthetic_hidden",
        "synthetic_date",
        "synthetic_datetime",
        "synthetic_real",
    )
    values = (
        "12345678009",
        " 3 ",
        " Synthetic Text ",
        date(2026, 9, 12),
        datetime(2026, 9, 12, 15, 30),
        None,
    )
    result = prepare(context, values, names)
    assert result.status == PreparationStatus.READY and len(result.prepared_values) == 4
    integer, text, day, moment = result.prepared_values
    assert integer.value_integer == 3 and text.value_text == "Synthetic Text"
    assert text.original_value == " Synthetic Text "
    assert day.value_date == "2026-09-12" and moment.value_date == "2026-09-12T15:30:00"
    for value in result.prepared_values:
        assert (
            sum(
                v is not None
                for v in (
                    value.value_number,
                    value.value_integer,
                    value.value_text,
                    value.value_date,
                )
            )
            == 1
        )


@pytest.mark.parametrize(
    "name,value",
    [
        ("synthetic_hidden", 3),
        ("synthetic_date", datetime(2026, 9, 12)),
        ("synthetic_integer", 3.2),
        ("synthetic_datetime", "09/12/2026"),
    ],
)
def test_conservative_type_failures_block_import(context, name, value):
    result = prepare(context, ("12345678009", value), ("stream_code", name))
    assert result.status == PreparationStatus.BLOCKED and result.prepared_values == ()
    assert "CONVERSION_FAILED" in {i.code for i in result.issues}


def test_excluded_formula_never_becomes_candidate(context):
    before = preview(context)
    last = before.mapped_values[-1]
    before = replace(
        before,
        mapped_values=(
            *before.mapped_values[:-1],
            replace(last, cell=replace(last.cell, is_formula=True, value="=SYNTHETIC_FORMULA")),
        ),
    )
    result = context.service.prepare_row(
        before, field_policy=ImportFieldPolicy(frozenset({"synthetic_real"}))
    )
    assert result.status == PreparationStatus.READY and result.prepared_values == ()
    assert "SYNTHETIC_FORMULA" not in str(asdict(result))


def test_identity_exclusion_has_no_alternate_value_path(context):
    result = prepare(context, field_policy=ImportFieldPolicy(frozenset({"stream_code"})))
    assert result.status == PreparationStatus.BLOCKED
    assert result.stream_code is None and result.core_data is None and not result.prepared_values


def test_unmapped_and_do_not_map_are_not_prepared(context):
    selected = mappings(context, (*DEFAULT_NAMES, "synthetic_hidden"))
    for status, method in (
        (MappingStatus.UNMAPPED, MappingMethod.NONE),
        (MappingStatus.DO_NOT_MAP, MappingMethod.USER),
    ):
        changed = (
            *selected[:-1],
            replace(selected[-1], dictionary_id=None, mapping_status=status, mapping_method=method),
        )
        before = context.preview.build_preview(
            [input_row((*DEFAULT_VALUES, "SYNTHETIC_UNMAPPED"))], changed
        )
        result = context.service.prepare(before.rows)
        assert result.summary.ready_rows == 1 and result.summary.prepared_value_count == 1
        assert "SYNTHETIC_UNMAPPED" not in str(asdict(result))
