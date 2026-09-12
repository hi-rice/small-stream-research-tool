"""Synthetic 사전·원본·임시 DB로 Preview를 검증한다. INSERT는 fixture에서만 수행한다."""

import json
import sqlite3
from contextlib import closing
from dataclasses import asdict, replace
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from small_stream_research_tool.database import connect_database, initialize_database
from small_stream_research_tool.models.column_mapping import (
    ColumnMappingDraft,
    MappingHeaderPart,
    MappingMethod,
    MappingStatus,
)
from small_stream_research_tool.models.excel import ExcelCell, ExcelRow
from small_stream_research_tool.models.import_preview import PreviewFieldPolicy, PreviewStatus
from small_stream_research_tool.models.import_preview_errors import (
    InvalidPreviewArgumentError,
    PreviewSourceChangedError,
    StreamLookupError,
)
from small_stream_research_tool.models.stream_code import CodeStatus, ComparisonStatus
from small_stream_research_tool.models.workspace import WorkspaceStep
from small_stream_research_tool.repositories.dictionary_repository import DictionaryRepository
from small_stream_research_tool.repositories.stream_lookup_repository import StreamLookupRepository
from small_stream_research_tool.services.dictionary_service import DictionaryService
from small_stream_research_tool.services.import_preview_service import ImportPreviewService
from small_stream_research_tool.services.workspace_service import WorkspaceService

NAMES = (
    "stream_code",
    "province_code",
    "city_county_code",
    "town_code",
    "stream_serial_no",
    "stream_name",
    "test_metric",
    "test_hidden",
    "test_masked",
)
VALUES = (
    "12345678009",
    "12",
    "345",
    "678",
    "009",
    "Synthetic Stream A",
    2.5,
    "SYNTHETIC_HIDDEN_VALUE",
    "SYNTHETIC_MASKED_VALUE",
)


@pytest.fixture
def context(tmp_path):
    path = tmp_path / "synthetic.sqlite3"
    initialize_database(path)
    with closing(connect_database(path)) as connection:
        dictionary = DictionaryService(DictionaryRepository(connection))
        category = dictionary.create_category("synthetic", "Synthetic Category")
        version = dictionary.create_version("synthetic-v1")
        items = {}
        for name in NAMES:
            item = dictionary.create_item(
                "Synthetic " + name,
                name,
                category.category_id,
                "TEXT",
                version.version_id,
                required=True,
            )
            dictionary.register_alias(item.dictionary_id, "Synthetic " + name)
            items[name] = item
        # 기존 하천은 비활성 상태여도 PK 존재 여부를 조회한다.
        connection.execute(
            """INSERT INTO small_stream
            (stream_code, province_code, city_county_code, town_code, stream_serial_no,
             stream_name, is_active, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, 0, ?, ?)""",
            (*VALUES[:6], "2020-01-01T00:00:00Z", "2020-01-01T00:00:00Z"),
        )
        lookup = StreamLookupRepository(connection)
        yield SimpleNamespace(
            connection=connection,
            dictionary=dictionary,
            items=items,
            lookup=lookup,
            service=ImportPreviewService(dictionary, lookup),
        )


def mappings(context, names=NAMES):
    return tuple(
        ColumnMappingDraft(
            index,
            get_column_letter(index),
            "Synthetic " + name,
            (MappingHeaderPart(1, "Synthetic " + name, None, None, None, None),),
            context.items[name].dictionary_id,
            MappingMethod.USER,
            MappingStatus.USER_MAPPED,
        )
        for index, name in enumerate(names, 1)
    )


def row(values=VALUES, source_row=2):
    return ExcelRow(
        source_row,
        tuple(
            ExcelCell(
                source_row,
                index,
                get_column_letter(index),
                value,
                "s" if isinstance(value, str) else "n",
                False,
                "General",
            )
            for index, value in enumerate(values, 1)
        ),
    )


def preview(context, values=VALUES, **options):
    return context.service.build_preview([row(values)], mappings(context), **options).rows[0]


@pytest.mark.parametrize(
    "source,components,status,effective",
    [
        ("12345678009", ("12", "345", "678", "009"), "EXISTING_STREAM", "12345678009"),
        ("12345678008", ("12", "345", "678", "008"), "NEW_STREAM", "12345678008"),
        (None, ("12", "345", "678", "009"), "EXISTING_STREAM", "12345678009"),
        ("12345678009", (None, None, None, None), "EXISTING_STREAM", "12345678009"),
        ("12345678008", ("12", "345", "678", "009"), "NEEDS_REVIEW", None),
        ("invalid", ("12", "345", "678", "009"), "NEEDS_REVIEW", None),
        (None, (None, None, None, None), "NEEDS_REVIEW", None),
        ("12345678009", ("12", "345", "678", 9), "NEEDS_REVIEW", None),
        ("12345678009", ("12", "345", "678", "X"), "NEEDS_REVIEW", None),
    ],
)
def test_identity_selection(context, source, components, status, effective):
    result = preview(context, (source, *components, *VALUES[5:]))
    assert result.status == status and result.stream_code == effective
    assert result.stream_code_validation.source_code == source
    if status == "NEEDS_REVIEW":
        assert result.existing_stream_found is None and any(i.blocking for i in result.issues)
    else:
        assert not any(i.blocking for i in result.issues)


def test_source_only_no_name_or_other_required_fields_not_blocked(context):
    result = context.service.build_preview(
        [row((VALUES[0],))], mappings(context, ("stream_code",))
    ).rows[0]
    assert result.status == PreviewStatus.EXISTING_STREAM and result.stream_name is None
    assert result.stream_code_validation.generated_code is None
    assert result.stream_code_validation.comparison_status == ComparisonStatus.NOT_COMPARABLE
    assert all(not issue.blocking for issue in result.issues)


def test_components_only_identity(context):
    names = NAMES[1:5]
    result = context.service.build_preview([row(VALUES[1:5])], mappings(context, names)).rows[0]
    assert result.status == PreviewStatus.EXISTING_STREAM
    assert result.stream_code_validation.source.status == CodeStatus.MISSING


def test_index_join_despite_duplicate_headers_and_shuffled_cells(context):
    source = row((VALUES[0], "Synthetic Stream B"))
    source = replace(source, cells=tuple(reversed(source.cells)))
    draft = mappings(context, ("stream_code", "stream_name"))
    draft = tuple(replace(m, source_header="Repeated Synthetic Header") for m in draft)
    result = context.service.build_preview([source], draft).rows[0]
    assert result.stream_code == VALUES[0] and result.stream_name == "Synthetic Stream B"
    assert result.source_row == 2
    assert [v.source_column_index for v in result.mapped_values] == [1, 2]


def test_auto_mapping_eligible(context):
    draft = tuple(
        replace(
            m, mapping_status=MappingStatus.AUTO_MAPPED, mapping_method=MappingMethod.AUTO_ALIAS
        )
        for m in mappings(context)
    )
    assert (
        context.service.build_preview([row()], draft).rows[0].status
        == PreviewStatus.EXISTING_STREAM
    )


@pytest.mark.parametrize(
    "state,method",
    [(MappingStatus.UNMAPPED, MappingMethod.NONE), (MappingStatus.DO_NOT_MAP, MappingMethod.USER)],
)
def test_unused_mapping_values_not_consumed(context, state, method):
    draft = list(mappings(context))
    draft[-1] = replace(draft[-1], mapping_status=state, mapping_method=method, dictionary_id=None)
    result = context.service.build_preview([row()], draft).rows[0]
    assert result.status == PreviewStatus.EXISTING_STREAM
    assert all(v.internal_name != "test_masked" for v in result.mapped_values)


def test_review_mapping_preserved_and_blocks(context):
    draft = list(mappings(context))
    draft[1] = replace(draft[1], mapping_status=MappingStatus.NEEDS_REVIEW)
    before = tuple(draft)
    result = context.service.build_preview([row()], draft).rows[0]
    assert result.status == PreviewStatus.NEEDS_REVIEW
    assert result.stream_code_validation.components[0].status == CodeStatus.MISSING
    assert any(i.code == "MAPPING_NEEDS_REVIEW" for i in result.issues)
    assert tuple(draft) == before


def test_dictionary_deactivated_after_mapping(context):
    draft = mappings(context)
    context.dictionary.deactivate_item(context.items["stream_code"].dictionary_id)
    result = context.service.build_preview([row()], draft).rows[0]
    assert result.status == PreviewStatus.NEEDS_REVIEW
    assert draft[0].mapping_status == MappingStatus.USER_MAPPED


@pytest.mark.parametrize("name", ["stream_code", "stream_name", "province_code", "test_metric"])
def test_duplicate_semantic_never_picks_first_or_last(context, name):
    draft = mappings(context, ("stream_code", name, name))
    result = context.service.build_preview([row((VALUES[0], "left", "right"))], draft).rows[0]
    assert result.status == PreviewStatus.NEEDS_REVIEW
    assert any(i.code == "DUPLICATE_SEMANTIC_MAPPING" for i in result.issues)
    if name == "stream_code":
        assert result.stream_code is None and result.stream_code_validation.source_code is None
    if name == "stream_name":
        assert result.stream_name is None


@pytest.mark.parametrize("field", [0, 4])
@pytest.mark.parametrize("kind", ["formula", "error"])
def test_code_formula_and_error_metadata(context, field, kind):
    source = row()
    cells = list(source.cells)
    # 정상 코드처럼 보이는 native 값이어도 metadata가 수식/오류이면 사용하지 않는다.
    cells[field] = replace(
        cells[field], is_formula=kind == "formula", value_type="f" if kind == "formula" else "e"
    )
    result = context.service.build_preview(
        [replace(source, cells=tuple(cells))], mappings(context)
    ).rows[0]
    assert result.status == PreviewStatus.NEEDS_REVIEW and result.stream_code is None
    assert any(i.code == "EXCEL_FORMULA_OR_ERROR" for i in result.issues)


def test_number_format_forwarded_without_duplicate_normalization(context):
    cells = list(row().cells)
    cells[4] = replace(cells[4], value=9, value_type="n", number_format="000")
    result = context.service.build_preview([ExcelRow(2, tuple(cells))], mappings(context)).rows[0]
    assert result.status == PreviewStatus.EXISTING_STREAM
    serial = result.stream_code_validation.components[-1]
    assert serial.raw_value == 9 and serial.normalized_value == "009"
    assert serial.normalization_steps == ("INTEGER_TO_TEXT", "ZERO_FORMAT_PADDING")


@pytest.mark.parametrize(
    "case", ["mismatch", "ambiguous", "invalid", "missing", "excluded", "duplicate"]
)
def test_unsafe_or_excluded_rows_never_lookup(context, monkeypatch, case):
    def forbidden(*args, **kwargs):
        pytest.fail("No existing stream lookup is allowed for this Preview")

    monkeypatch.setattr(context.lookup, "exists_by_stream_code", forbidden)
    values, draft, options = list(VALUES), mappings(context), {}
    if case == "mismatch":
        values[0] = "12345678008"
    elif case == "ambiguous":
        values[4] = 9
    elif case == "invalid":
        values[0] = "X"
    elif case == "missing":
        values[:5] = [None] * 5
    elif case == "excluded":
        options["excluded_rows"] = {2}
    else:
        draft = (draft[0], replace(draft[1], dictionary_id=draft[0].dictionary_id))
    result = context.service.build_preview([row(values)], draft, **options).rows[0]
    assert result.existing_stream_found is None


def test_blank_partial_trailing_excluded_and_summary(context):
    values = [
        row(),
        row(("12345678008", "12", "345", "678", "008", *VALUES[5:]), 3),
        row((None,) * 9, 4),
        row((None,) * 8 + ("synthetic trailing",), 5),
        row(VALUES, 6),
    ]
    result = context.service.build_preview(values, mappings(context), excluded_rows={6})
    assert [r.source_row for r in result.rows] == [2, 3, 5, 6]
    assert [r.status for r in result.rows] == [
        PreviewStatus.EXISTING_STREAM,
        PreviewStatus.NEW_STREAM,
        PreviewStatus.NEEDS_REVIEW,
        PreviewStatus.EXCLUDED,
    ]
    assert asdict(result.summary) == dict(
        total_rows=4, new_stream=1, existing_stream=1, needs_review=1, excluded=1
    )
    assert (
        context.service.build_preview(
            [values[2]], mappings(context), include_blank=True
        ).summary.needs_review
        == 1
    )
    assert (
        context.service.build_preview(
            [values[2]], mappings(context), excluded_rows={4}
        ).summary.excluded
        == 1
    )


def test_display_exclusion_masking_internal_names_only(context, caplog):
    policy = PreviewFieldPolicy(
        frozenset({"test_hidden"}), frozenset({"test_masked", "stream_name"})
    )
    result = preview(context, field_policy=policy)
    safe = result.to_display()
    payload = json.dumps(asdict(safe), default=str)
    assert "SYNTHETIC_HIDDEN_VALUE" not in payload and "SYNTHETIC_MASKED_VALUE" not in payload
    assert "Synthetic Stream A" not in payload
    assert safe.stream_name == "[MASKED]"
    assert safe.effective_code == VALUES[0] and safe.source_code == safe.generated_code == VALUES[0]
    assert any(v.internal_name == "test_metric" and v.value == 2.5 for v in safe.display_values)
    assert all(v.internal_name != "test_hidden" for v in safe.display_values)
    assert result.mapped_values[-2].cell.value == "SYNTHETIC_HIDDEN_VALUE"
    assert "SYNTHETIC_HIDDEN_VALUE" not in repr(result) and caplog.records == []
    # raw header를 정책 key로 전달해도 internal_name과 같지 않으면 해당 필드를 숨기지 않는다.
    wrong_policy = PreviewFieldPolicy(frozenset({"Synthetic test_hidden"}))
    assert any(
        v.internal_name == "test_hidden"
        for v in preview(context, field_policy=wrong_policy).display_values
    )
    assert PreviewFieldPolicy().excluded_internal_names == frozenset()


def test_excluded_stream_name_has_no_secondary_display_leak(context):
    policy = PreviewFieldPolicy(frozenset({"stream_name"}))
    result = preview(context, field_policy=policy)
    assert result.stream_name is None and "Synthetic Stream A" not in str(
        asdict(result.to_display())
    )


@pytest.mark.parametrize(
    "policy",
    [
        PreviewFieldPolicy(frozenset({"stream_code"})),
        PreviewFieldPolicy(masked_internal_names=frozenset({"province_code"})),
        PreviewFieldPolicy({"test_hidden"}),
        PreviewFieldPolicy(frozenset({""})),
        "invalid",
    ],
)
def test_invalid_policy(context, policy):
    with pytest.raises(InvalidPreviewArgumentError):
        preview(context, field_policy=policy)


def test_select_only_no_commit_or_db_mutation(context):
    connection = context.connection
    before = connection.total_changes
    allowed = {sqlite3.SQLITE_SELECT, sqlite3.SQLITE_READ, sqlite3.SQLITE_FUNCTION}
    actions = []

    def authorize(action, *args):
        actions.append(action)
        return sqlite3.SQLITE_OK if action in allowed else sqlite3.SQLITE_DENY

    connection.set_authorizer(authorize)
    try:
        result = preview(context)
        assert result.status == PreviewStatus.EXISTING_STREAM
        for table in (
            "source_file",
            "import_history",
            "import_sheet",
            "import_column_mapping",
            "characteristic_value",
            "data_quality_issue",
            "record_history",
        ):
            assert connection.execute(f"SELECT count(*) FROM {table}").fetchone() == (0,)
        assert connection.total_changes == before and set(actions) <= allowed
    finally:
        connection.set_authorizer(None)
    connection.execute("BEGIN")
    try:
        preview(context)
        assert connection.in_transaction
    finally:
        connection.rollback()


@pytest.mark.parametrize(
    "value", [None, 12345678009, "123", "1234567800A", "１２３４５６７８００９"]
)
def test_lookup_rejects_unvalidated_code(context, value):
    with pytest.raises(InvalidPreviewArgumentError):
        context.lookup.exists_by_stream_code(value)


def test_missing_db_schema_fails_without_fallback(tmp_path):
    with closing(connect_database(tmp_path / "uninitialized.sqlite3")) as connection:
        repository = StreamLookupRepository(connection)
        with pytest.raises(StreamLookupError):
            repository.exists_by_stream_code(VALUES[0])
        assert connection.execute("SELECT count(*) FROM sqlite_schema").fetchone() == (0,)


def test_determinism_and_immutability(context):
    source, draft = row(), mappings(context)
    before = (asdict(source), tuple(asdict(m) for m in draft))
    first = context.service.build_preview([source], draft)
    assert first == context.service.build_preview([source], draft)
    assert before == (asdict(source), tuple(asdict(m) for m in draft))
    assert "UPDATED" not in [status.value for status in PreviewStatus]


def test_workspace_preview_step_and_rebuild(context, tmp_path):
    path = tmp_path / "synthetic.xlsx"
    workbook = Workbook()
    workbook.active.append(["Synthetic " + name for name in NAMES])
    workbook.active.append(VALUES)
    workbook.save(path)
    workbook.close()
    workspace = WorkspaceService(tmp_path / "workspace")
    draft = workspace.create_workspace(
        current_user_id=1,
        source_file_path=path,
        selected_sheet_name="Sheet",
        header_start_row=1,
        header_end_row=1,
        data_start_row=2,
        column_mappings=mappings(context),
        current_step=WorkspaceStep.PREVIEW,
    )
    for step in (WorkspaceStep.MAPPING, WorkspaceStep.CODE_VALIDATION, WorkspaceStep.PREVIEW):
        saved = workspace.save_workspace(replace(draft, current_step=step), 1)
        assert workspace.load_workspace(1) == saved and saved.workspace_version == 1
    before = (tmp_path / "workspace" / "user_1.json").read_bytes()
    policy = PreviewFieldPolicy(frozenset({"test_hidden"}))
    result = context.service.rebuild_from_workspace(workspace, 1, field_policy=policy)
    assert result == context.service.build_preview([row()], mappings(context), field_policy=policy)
    assert (tmp_path / "workspace" / "user_1.json").read_bytes() == before
    data = json.loads(before)
    assert "rows" not in data and "preview" not in data and b"SYNTHETIC_HIDDEN_VALUE" not in before
    assert context.service.rebuild_from_workspace(workspace, 2) is None


def test_rebuild_source_change_guard(context, tmp_path, monkeypatch):
    import small_stream_research_tool.services.import_preview_service as module

    path = tmp_path / "synthetic.xlsx"
    workbook = Workbook()
    workbook.active.append(["Synthetic stream_code"])
    workbook.active.append([VALUES[0]])
    workbook.save(path)
    workbook.close()
    workspace = WorkspaceService(tmp_path / "workspace")
    draft = workspace.create_workspace(
        current_user_id=1,
        source_file_path=path,
        selected_sheet_name="Sheet",
        header_start_row=1,
        header_end_row=1,
        data_start_row=2,
        column_mappings=mappings(context, ("stream_code",)),
        current_step=WorkspaceStep.PREVIEW,
    )
    workspace.save_workspace(draft, 1)
    monkeypatch.setattr(module, "file_sha256", lambda path: "0" * 64)
    with pytest.raises(PreviewSourceChangedError):
        context.service.rebuild_from_workspace(workspace, 1)


@pytest.mark.parametrize("case", ["row_zero", "row_bool", "cell_row", "duplicate_column", "letter"])
def test_invalid_source_coordinates(context, case):
    source = row()
    if case == "row_zero":
        source = replace(source, row_index=0)
    elif case == "row_bool":
        source = replace(source, row_index=True)
    elif case == "cell_row":
        source = replace(source, cells=(replace(source.cells[0], row_index=99),))
    elif case == "duplicate_column":
        source = replace(source, cells=(source.cells[0], source.cells[0]))
    else:
        source = replace(source, cells=(replace(source.cells[0], column_letter="Z"),))
    with pytest.raises(InvalidPreviewArgumentError):
        context.service.build_preview([source], mappings(context))


def test_missing_source_column_is_reviewed(context):
    result = context.service.build_preview([row((VALUES[0],))], mappings(context)).rows[0]
    assert result.status == PreviewStatus.NEEDS_REVIEW
    assert any(issue.code == "SOURCE_COLUMN_MISSING" for issue in result.issues)


def test_iterator_is_lazy_and_duplicate_rows_rejected(context, monkeypatch):
    calls = []
    original = context.lookup.exists_by_stream_code

    def observe(code):
        calls.append(code)
        return original(code)

    monkeypatch.setattr(context.lookup, "exists_by_stream_code", observe)
    iterator = context.service.iter_preview_rows(iter([row(), row()]), mappings(context))
    assert iter(iterator) is iterator and calls == []
    assert next(iterator).status == PreviewStatus.EXISTING_STREAM and len(calls) == 1
    with pytest.raises(InvalidPreviewArgumentError):
        next(iterator)
