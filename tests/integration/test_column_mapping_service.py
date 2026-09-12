"""실제 사전 Service와 임시 SQLite를 사용하되 매핑 작업에서는 DB를 쓰지 않는다."""

from contextlib import closing
from dataclasses import replace

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from small_stream_research_tool.database import connect_database, initialize_database
from small_stream_research_tool.models.column_mapping import MappingMethod, MappingStatus
from small_stream_research_tool.models.excel import ExcelCell, ExcelColumn, ExcelHeaderPart
from small_stream_research_tool.models.mapping_errors import (
    InvalidMappingError,
    MappingTargetUnavailableError,
)
from small_stream_research_tool.models.workspace import WorkspaceStep
from small_stream_research_tool.repositories.dictionary_repository import DictionaryRepository
from small_stream_research_tool.services.column_mapping_service import (
    ColumnMappingService,
    mapping_summary,
    validate_mapping,
)
from small_stream_research_tool.services.dictionary_service import DictionaryService
from small_stream_research_tool.services.excel_reader import ExcelReader
from small_stream_research_tool.services.workspace_service import WorkspaceService


def column(index=1, header="Synthetic Area"):
    cell = ExcelCell(1, index, get_column_letter(index), header, "s", False, "General")
    return ExcelColumn(index, cell.column_letter, (ExcelHeaderPart(cell, None, None),), header)


@pytest.fixture
def context(tmp_path):
    db = tmp_path / "synthetic.sqlite3"
    initialize_database(db)
    with closing(connect_database(db)) as connection:
        dictionary = DictionaryService(DictionaryRepository(connection))
        category = dictionary.create_category("synthetic", "Synthetic Category")
        version = dictionary.create_version("synthetic-v1")
        unit = dictionary.create_unit("Synthetic Unit", "synthetic-unit")
        first = dictionary.create_item(
            "Synthetic Area",
            "test_area",
            category.category_id,
            "REAL",
            version.version_id,
            unit_id=unit.unit_id,
        )
        second = dictionary.create_item(
            "Synthetic Other", "test_other", category.category_id, "TEXT", version.version_id
        )
        alias = dictionary.register_alias(first.dictionary_id, "Synthetic Area")
        yield connection, dictionary, ColumnMappingService(dictionary), first, second, alias


def test_initial_mapping_provenance_and_readonly(context):
    connection, _, service, first, _, _ = context
    before = connection.total_changes
    source = (column(1), column(5), column(9, ""), column(10, "unknown"))
    result = service.build_initial_mappings(source)
    assert [m.mapping_status for m in result] == [MappingStatus.AUTO_MAPPED] * 2 + [
        MappingStatus.UNMAPPED
    ] * 2
    assert [m.source_column_index for m in result] == [1, 5, 9, 10]
    assert result[0].dictionary_id == result[1].dictionary_id == first.dictionary_id
    assert result[0].source_header_parts[0].text == source[0].header_parts[0].cell.value
    assert result[1].source_column_letter == "E"
    assert connection.total_changes == before
    for table in ("source_file", "import_history", "import_sheet", "import_column_mapping"):
        assert connection.execute(f"SELECT count(*) FROM {table}").fetchone() == (0,)


@pytest.mark.parametrize(
    "header,matched",
    [
        (" Synthetic\t Area ", True),
        ("SyntheticArea", False),
        ("Synthetic", False),
        ("Synthetic Area(km²)", False),
        (" \t", False),
    ],
)
def test_normalization_only_registered_aliases(context, header, matched):
    _, _, service, _, _, _ = context
    result = service.build_initial_mappings([column(header=header)])[0]
    assert (result.mapping_status == MappingStatus.AUTO_MAPPED) == matched
    assert result.source_header == header


def test_exact_scope_fallback_and_inactive_shadow(context):
    _, dictionary, service, first, second, _ = context
    specific = dictionary.register_alias(
        second.dictionary_id, "Synthetic Area", source_scope="scope_alpha"
    )
    exact = service.build_initial_mappings([column()], "scope_alpha")[0]
    fallback = service.build_initial_mappings([column()], "scope_beta")[0]
    assert exact.dictionary_id == second.dictionary_id
    assert fallback.dictionary_id == first.dictionary_id
    dictionary.deactivate_alias(specific.alias_id)
    assert service.build_initial_mappings([column()], "scope_alpha")[0].dictionary_id is None


def test_manual_clear_exclude_no_alias_insertion(context):
    connection, _, service, _, second, _ = context
    original = service.build_initial_mappings([column()])[0]
    before = connection.total_changes
    user = service.set_user_mapping(original, second.dictionary_id)
    assert (user.mapping_method, user.mapping_status) == (
        MappingMethod.USER,
        MappingStatus.USER_MAPPED,
    )
    assert (
        user.dictionary_id == second.dictionary_id and original.dictionary_id != user.dictionary_id
    )
    cleared = service.clear_mapping(user)
    excluded = service.mark_do_not_map(user)
    assert cleared.mapping_status == MappingStatus.UNMAPPED and cleared.dictionary_id is None
    assert excluded.mapping_status == MappingStatus.DO_NOT_MAP and excluded.dictionary_id is None
    assert (
        cleared.source_header_parts == excluded.source_header_parts == original.source_header_parts
    )
    assert connection.total_changes == before


@pytest.mark.parametrize("target_state", ["missing", "inactive", "category", "unit", "deprecated"])
def test_manual_target_unavailable(context, target_state):
    _, dictionary, service, first, _, _ = context
    mapping = service.build_initial_mappings([column()])[0]
    target = first.dictionary_id
    if target_state == "missing":
        target = 999999
    elif target_state == "inactive":
        dictionary.deactivate_item(target)
    elif target_state == "category":
        dictionary.deactivate_category(first.category_id)
    elif target_state == "unit":
        dictionary.deactivate_unit(first.unit_id)
    else:
        version = dictionary.create_version("synthetic-v2")
        dictionary.deprecate_item(target, version.version_id)
    with pytest.raises(MappingTargetUnavailableError):
        service.set_user_mapping(mapping, target)


def test_refresh_preserves_user_choices_and_summary(context):
    _, dictionary, service, first, second, alias = context
    auto, unknown, manual, excluded = service.build_initial_mappings(
        [column(1), column(2, "Synthetic New"), column(3), column(4)]
    )
    manual = service.set_user_mapping(manual, second.dictionary_id)
    excluded = service.mark_do_not_map(excluded)
    review = replace(
        auto,
        source_column_index=5,
        source_column_letter="E",
        mapping_status=MappingStatus.NEEDS_REVIEW,
    )
    dictionary.deactivate_alias(alias.alias_id)
    dictionary.register_alias(first.dictionary_id, "Synthetic New")
    result = service.refresh_auto_mappings((auto, unknown, manual, excluded, review))
    assert result[0].mapping_status == MappingStatus.UNMAPPED
    assert result[1].mapping_status == MappingStatus.AUTO_MAPPED
    assert result[2:] == (manual, excluded, review)
    summary = mapping_summary(result)
    assert (
        summary.total_columns,
        summary.auto_mapped,
        summary.user_mapped,
        summary.unmapped,
        summary.do_not_map,
        summary.needs_review,
    ) == (5, 1, 1, 1, 1, 1)


@pytest.mark.parametrize(
    "changes",
    [
        {"source_column_index": 0},
        {"source_column_index": True},
        {"source_column_index": 16385},
        {"source_column_letter": "Z"},
        {"dictionary_id": None},
        {"dictionary_id": True},
        {"mapping_method": MappingMethod.NONE},
        {"mapping_status": "AUTO_MAPPED"},
        {"mapping_status": MappingStatus.DO_NOT_MAP},
        {"mapping_status": MappingStatus.UNMAPPED},
        {"source_header": None},
        {"source_scope": ""},
    ],
)
def test_invalid_mapping_combinations(context, changes):
    _, _, service, _, _, _ = context
    mapping = service.build_initial_mappings([column()])[0]
    with pytest.raises(InvalidMappingError):
        validate_mapping(replace(mapping, **changes))


def test_duplicate_index_rejected(context):
    _, _, service, _, _, _ = context
    with pytest.raises(InvalidMappingError):
        service.build_initial_mappings([column(), column()])


@pytest.mark.parametrize("change", ["none", "inactive", "missing", "alias_inactive"])
def test_resume_revalidates_without_replacing_id_or_saved_json(context, tmp_path, change):
    connection, dictionary, service, first, _, alias = context
    path = tmp_path / "synthetic.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Synthetic Area"
    sheet["A2"] = "SYNTHETIC_ROW_VALUE_NEVER_PERSIST"
    workbook.save(path)
    workbook.close()
    with ExcelReader(path) as reader:
        mappings = service.build_initial_mappings(
            reader.read_columns("Sheet", header_start_row=1, header_end_row=1)
        )
    if change == "missing":
        mappings = (
            replace(
                mappings[0],
                dictionary_id=999999,
                mapping_method=MappingMethod.USER,
                mapping_status=MappingStatus.USER_MAPPED,
            ),
        )
    workspace = WorkspaceService(tmp_path / "workspace")
    draft = workspace.create_workspace(
        current_user_id=1,
        source_file_path=path,
        selected_sheet_name="Sheet",
        header_start_row=1,
        header_end_row=1,
        data_start_row=2,
        column_mappings=mappings,
        current_step=WorkspaceStep.MAPPING,
    )
    workspace.save_workspace(draft, 1)
    stored = (tmp_path / "workspace" / "user_1.json").read_bytes()
    assert b"SYNTHETIC_ROW_VALUE_NEVER_PERSIST" not in stored
    if change == "inactive":
        dictionary.deactivate_item(first.dictionary_id)
    if change == "alias_inactive":
        dictionary.deactivate_alias(alias.alias_id)
    before = connection.total_changes
    result = service.resume_workspace(WorkspaceService(tmp_path / "workspace"), 1)
    assert result.original.column_mappings == mappings
    assert result.resumed.column_mappings[0].dictionary_id == mappings[0].dictionary_id
    expected = MappingStatus.AUTO_MAPPED if change == "none" else MappingStatus.NEEDS_REVIEW
    assert result.resumed.column_mappings[0].mapping_status == expected
    assert (tmp_path / "workspace" / "user_1.json").read_bytes() == stored
    assert connection.total_changes == before


def test_merged_header_structure_lookup(context, tmp_path):
    _, dictionary, service, first, _, _ = context
    dictionary.register_alias(first.dictionary_id, "Synthetic Group | Left")
    path = tmp_path / "merged.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Synthetic Group"
    sheet["A2"], sheet["B2"] = "Left", "Right"
    workbook.save(path)
    workbook.close()
    with ExcelReader(path) as reader:
        columns = reader.read_columns("Sheet", header_start_row=1, header_end_row=2)
        mappings = service.build_initial_mappings(columns)
    assert mappings[0].dictionary_id == first.dictionary_id
    assert mappings[1].mapping_status == MappingStatus.UNMAPPED
    part = mappings[1].source_header_parts[0]
    assert part.text is None and part.anchor_text == "Synthetic Group"
    assert (part.row_index, part.anchor_row, part.anchor_column, part.merged_range) == (
        1,
        1,
        1,
        "A1:B1",
    )
