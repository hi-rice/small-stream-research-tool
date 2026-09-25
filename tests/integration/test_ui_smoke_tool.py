"""개발 전용 UI smoke DB의 격리·합성 상태·Phase 8 연계를 검증한다."""

import time
import uuid
from contextlib import closing
from subprocess import run

import pytest
import tools.ui_smoke as ui_smoke
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication
from tools.ui_smoke import (
    PHASE10D_CHECKLIST,
    PHASE10D_HEADERS,
    PHASE10D_ROWS,
    PHASE10D_SHEET,
    ROOT,
    create_database,
    create_phase10d_smoke,
    finalize_after_registration,
    phase10d_workspace_path,
    phase10d_xlsx_path,
    smoke_path,
)

from small_stream_research_tool.app.controller import GuiSession
from small_stream_research_tool.config.settings import get_app_paths
from small_stream_research_tool.database import connect_database
from small_stream_research_tool.database.connection import transaction
from small_stream_research_tool.models.errors import AuthenticationError
from small_stream_research_tool.repositories.user_repository import UserRepository
from small_stream_research_tool.services.auth_service import AuthService
from small_stream_research_tool.ui.main_window import MainWindow


@pytest.fixture
def smoke_database():
    path = ROOT / "build" / f"test_ui_smoke_{uuid.uuid4().hex}.db"
    try:
        yield create_database(path)
    finally:
        path.unlink(missing_ok=True)


@pytest.fixture
def phase10d_artifacts():
    token = uuid.uuid4().hex
    db = ROOT / "build" / f"phase10d_ui_smoke_{token}.db"
    xlsx = ROOT / "build" / f"phase10d_ui_smoke_{token}.xlsx"
    try:
        yield create_phase10d_smoke(db, xlsx)
    finally:
        db.unlink(missing_ok=True)
        xlsx.unlink(missing_ok=True)


@pytest.fixture(scope="session")
def app():
    return QApplication.instance() or QApplication([])


def wait_for(app, condition, seconds=8):
    end = time.monotonic() + seconds
    while time.monotonic() < end:
        app.processEvents()
        if condition():
            return
        time.sleep(0.02)
    pytest.fail("UI smoke 화면 전환이 완료되지 않았습니다.")


def test_path_guard_refuses_default_outside_and_overwrite(smoke_database):
    default = get_app_paths().database_dir / "research.sqlite3"
    with pytest.raises(ValueError):
        smoke_path(default)
    with pytest.raises(ValueError):
        smoke_path(ROOT / "ui_smoke.db")
    with pytest.raises(FileExistsError):
        create_database(smoke_database)


def test_phase10d_path_guards_refuse_outside_and_overwrite(phase10d_artifacts, tmp_path):
    db, xlsx = phase10d_artifacts
    with pytest.raises(ValueError):
        phase10d_xlsx_path(tmp_path / "phase10d_ui_smoke.xlsx")
    with pytest.raises(ValueError):
        create_phase10d_smoke(tmp_path / "phase10d_ui_smoke.db", xlsx)
    with pytest.raises(FileExistsError):
        create_phase10d_smoke(db, ROOT / "build" / "phase10d_ui_smoke_new.xlsx")
    with pytest.raises(FileExistsError):
        create_phase10d_smoke(
            ROOT / "build" / "phase10d_ui_smoke_new.db",
            xlsx,
        )


def test_phase10d_artifacts_use_official_v2_core_and_synthetic_workbook(phase10d_artifacts):
    db, xlsx = phase10d_artifacts
    with closing(connect_database(db)) as connection:
        assert connection.execute(
            "SELECT version FROM dictionary_version WHERE is_current=1"
        ).fetchone() == ("research-dictionary-v2",)
        assert connection.execute("SELECT count(*) FROM data_dictionary").fetchone() == (76,)
        assert connection.execute("SELECT count(*) FROM app_user").fetchone() == (0,)
        assert connection.execute("SELECT count(*) FROM import_history").fetchone() == (0,)
        assert connection.execute("SELECT count(*) FROM characteristic_value").fetchone() == (0,)
        assert connection.execute("SELECT count(*) FROM data_quality_issue").fetchone() == (0,)
        assert connection.execute("SELECT count(*) FROM stream_characteristic").fetchone() == (0,)
        assert connection.execute(
            "SELECT stream_code,stream_name FROM small_stream"
        ).fetchall() == [("09876543001", "PHASE10D 합성 기존천")]
        aliases = {
            row[0]
            for row in connection.execute(
                "SELECT alias_name FROM column_alias WHERE source_scope='NATIONAL_2024'"
            )
        }
        assert set(PHASE10D_HEADERS[:-1]) <= aliases
        assert sum(header in aliases for header in PHASE10D_HEADERS) == 9
        assert tuple(header for header in PHASE10D_HEADERS if header not in aliases) == (
            "검토메모",
        )

    book = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        sheet = book[PHASE10D_SHEET]
        assert sheet.max_row == len(PHASE10D_ROWS) + 1
        assert tuple(cell.value for cell in sheet[1]) == PHASE10D_HEADERS
        rows = tuple(sheet.iter_rows(min_row=2, values_only=True))
        assert rows == PHASE10D_ROWS
        assert all(row[0].startswith("0") and len(row[0]) == 11 for row in rows)
        assert all("PHASE10D 합성" in row[5] for row in rows)
    finally:
        book.close()

    workspace = phase10d_workspace_path(db)
    assert not workspace.exists()
    assert not workspace.resolve().is_relative_to(ROOT.resolve())
    for artifact in (db, xlsx):
        result = run(
            ["git", "check-ignore", "--quiet", str(artifact.relative_to(ROOT))],
            cwd=ROOT,
            check=False,
        )
        assert result.returncode == 0


def test_phase10d_checklist_requires_scope_before_mapping_review():
    instructions = "\n".join(PHASE10D_CHECKLIST)
    assert "2024 전국 연구자료" in instructions
    assert "자료 범위 적용" in instructions
    assert "자동 매핑 9개" in instructions
    assert "검토메모" in instructions
    assert "전체 4행" in instructions


def test_phase10d_run_uses_isolated_workspace_without_legacy_finalizer(
    phase10d_artifacts, monkeypatch
):
    db, _xlsx = phase10d_artifacts
    captured = {}

    class Auth:
        @staticmethod
        def needs_initial_user_setup():
            return False

    class Controller:
        def __init__(self, db_path, workspace_dir):
            captured["db"] = db_path
            captured["workspace"] = workspace_dir
            self.auth_service = Auth()

        def start(self):
            captured["started"] = True

        def close(self):
            pass

    monkeypatch.setattr(ui_smoke, "ApplicationController", Controller)
    monkeypatch.setattr(
        ui_smoke,
        "finalize_after_registration",
        lambda _path: pytest.fail("Phase 9C finalizer must not run for Phase 10D"),
    )
    monkeypatch.setattr(QApplication, "exec", lambda self: 0)
    assert ui_smoke.run_gui(db) == 0
    assert captured == {
        "db": db,
        "workspace": phase10d_workspace_path(db),
        "started": True,
    }


def test_created_database_is_synthetic_unregistered_and_bootstrapped(smoke_database):
    with closing(connect_database(smoke_database)) as connection:
        assert connection.execute("SELECT count(*) FROM app_user").fetchone() == (0,)
        assert connection.execute("SELECT count(*) FROM small_stream").fetchone() == (5,)
        assert connection.execute("SELECT count(*) FROM data_dictionary").fetchone() == (76,)
        assert connection.execute(
            "SELECT count(*) FROM dictionary_version "
            "WHERE version='research-dictionary-v1' AND is_current=1"
        ).fetchone() == (1,)
        names = {row[0] for row in connection.execute("SELECT stream_name FROM small_stream")}
        assert names == {"테스트천A", "테스트천B", "테스트천C", "테스트천D", "테스트천E"}
        assert connection.execute("SELECT count(*) FROM characteristic_value").fetchone()[0] >= 100
        assert connection.execute("SELECT original_path FROM source_file").fetchone() == (
            "synthetic://ui-smoke/synthetic_ui_smoke.xlsx",
        )


def test_registration_finalizer_uses_services_and_builds_ui_states(smoke_database):
    with closing(connect_database(smoke_database)) as connection:
        with transaction(connection):
            UserRepository(connection).create_user(
                login_id="synthetic_actor",
                password_hash="NOT_A_LOGIN_CREDENTIAL",
                display_name="합성 UI 검수자",
                department="합성 부서",
                role=None,
                timestamp="2026-09-17T00:01:00Z",
            )
    assert finalize_after_registration(smoke_database)
    assert finalize_after_registration(smoke_database)
    with closing(connect_database(smoke_database)) as connection:
        representatives = dict(
            connection.execute(
                "SELECT stream_code,count(*) FROM characteristic_value "
                "WHERE is_active=1 AND is_representative=1 GROUP BY stream_code"
            )
        )
        caches = dict(
            connection.execute(
                "SELECT stream_code,count(*) FROM stream_characteristic GROUP BY stream_code"
            )
        )
        assert representatives["99000000001"] == 70
        assert 0 < representatives["99000000004"] < 70
        assert representatives["99000000005"] > 0
        assert caches.get("99000000005", 0) == 0
        assert connection.execute(
            "SELECT count(*) FROM characteristic_value WHERE source_type='USER_CORRECTION'"
        ).fetchone() == (1,)
        assert connection.execute(
            "SELECT count(*) FROM characteristic_value "
            "WHERE source_type='USER_CORRECTION' AND is_representative=1"
        ).fetchone() == (1,)
        assert connection.execute(
            "SELECT count(*) FROM record_history WHERE change_type='CORRECTION'"
        ).fetchone() == (1,)
        states = set(
            connection.execute(
                "SELECT severity,review_status FROM data_quality_issue WHERE is_active=1"
            )
        )
        assert ("WARNING", "CONFIRMED") in states
        assert ("ERROR", "IN_REVIEW") in states


def test_smoke_registration_and_login_use_same_database_and_password(smoke_database):
    password = "synthetic smoke passphrase"
    with closing(connect_database(smoke_database)) as connection:
        service = AuthService(UserRepository(connection))
        assert service.needs_initial_user_setup()
        created = service.create_user("smoke.researcher", password, "합성 검수자", department=None)
        assert created.department is None
    assert finalize_after_registration(smoke_database)
    with closing(connect_database(smoke_database)) as connection:
        service = AuthService(UserRepository(connection))
        authenticated = service.authenticate(" SMOKE.RESEARCHER ", password)
        assert authenticated.user_id == created.user_id
        with pytest.raises(AuthenticationError):
            service.authenticate("smoke.researcher", "wrong synthetic passphrase")


def test_offscreen_smoke_browse_detail_tabs_back_and_logout(smoke_database, app):
    with closing(connect_database(smoke_database)) as connection:
        service = AuthService(UserRepository(connection))
        user = service.create_user("smoke.researcher", "synthetic smoke passphrase", "합성 검수자")
    assert finalize_after_registration(smoke_database)
    window = MainWindow(smoke_database, GuiSession(user.user_id, user.display_name, None, None))
    logout = []
    window.logout_requested.connect(lambda: logout.append(True))
    try:
        wait_for(app, lambda: window.stream_list.model.rowCount() == 5)
        view = window.stream_list
        row = next(
            index
            for index in range(view.model.rowCount())
            if view.model.stream_code(index) == "99000000001"
        )
        view.table.selectRow(row)
        wait_for(app, lambda: view.summary_name.text() == "테스트천A")
        view.detail_button.click()
        detail = window.stream_detail
        wait_for(app, lambda: detail.status.text() == "읽기 전용 상세정보")
        assert detail.category.count() == 5
        for index in range(detail.category.count()):
            detail.category.setCurrentIndex(index)
            app.processEvents()
            assert detail.model.rowCount() > 0
        detail.table.selectRow(0)
        for index in range(detail.detail_tabs.count()):
            detail.detail_tabs.setCurrentIndex(index)
            app.processEvents()
        detail.back.click()
        assert window.stack.currentWidget() is view
        window.logout_button.click()
        assert logout == [True]
    finally:
        window.close()
        app.processEvents()
