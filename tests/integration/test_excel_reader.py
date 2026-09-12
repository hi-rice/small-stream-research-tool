"""실제 연구자료 없이 임시 synthetic xlsx로 Reader 계약을 검증한다."""

import hashlib
import sqlite3
from dataclasses import FrozenInstanceError
from datetime import date, datetime, time, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from small_stream_research_tool.models.excel import ExcelFormula
from small_stream_research_tool.models.excel_errors import (
    ExcelFileNotFoundError,
    ExcelReaderClosedError,
    ExcelSheetNotFoundError,
    ExcelWorkbookReadError,
    InvalidDataRangeError,
    InvalidHeaderRangeError,
    UnsupportedExcelFormatError,
    UnsupportedExcelSheetError,
)
from small_stream_research_tool.services.excel_reader import ExcelReader


@pytest.fixture
def workbook_path(tmp_path):
    path = tmp_path / "가상 자료.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SingleHeader"
    sheet.append(["TEXT", "INTEGER", "REAL", "DATE", None, "formula", "duplicate", "duplicate"])
    sheet.append(["001", 1, 3.5, date(2020, 1, 2), None, "=B2+C2", "3.50", "#N/A"])
    sheet["B2"].number_format = "000"
    sheet.append([False, 0, " ", datetime(2020, 1, 2, 3, 4, 5), time(1, 2), None, None, None])
    sheet["H5"] = "synthetic trailing value"
    multi = workbook.create_sheet("MultiHeader")
    multi["A1"] = "Synthetic title"
    multi.merge_cells("A2:C2")
    multi["A2"] = "Synthetic group"
    multi["A3"], multi["B3"], multi["C3"] = "left", "middle", "right"
    multi.merge_cells("D2:D4")
    multi["D2"] = "Vertical"
    multi["A4"], multi["B4"], multi["C4"] = "leaf", 7, date(2020, 2, 3)
    multi["E4"] = "Unmerged"
    multi["G6"] = "synthetic last column"
    hidden = workbook.create_sheet("HiddenSheet")
    hidden.sheet_state = "hidden"
    hidden.append(["Hidden header"])
    hidden.append(["Hidden data"])
    workbook.save(path)
    workbook.close()
    return path


def columns(reader, sheet="SingleHeader", start=1, end=1):
    return reader.read_columns(sheet, header_start_row=start, header_end_row=end)


def rows(reader, **kwargs):
    return reader.iter_rows("SingleHeader", header_start_row=1, header_end_row=1, **kwargs)


def test_workbook_and_sheet_metadata(workbook_path):
    with ExcelReader(workbook_path) as reader:
        info = reader.workbook_info
        assert info.file_path == workbook_path.absolute()
        assert info.sheet_names == ("SingleHeader", "MultiHeader", "HiddenSheet")
        assert info.sheet_count == 3
        sheet = reader.sheet_info("SingleHeader")
        assert (sheet.index, sheet.max_row, sheet.max_column, sheet.state) == (1, 5, 8, "visible")
        assert reader.sheet_info("HiddenSheet").state == "hidden"
        assert reader.sheet_info("HiddenSheet").index == 3
        assert reader.sheet_info("MultiHeader").merged_ranges == ("A2:C2", "D2:D4")
        assert reader.read_preview("HiddenSheet", header_start_row=1, header_end_row=1).rows
        with pytest.raises(FrozenInstanceError):
            info.file_path = Path("changed")


@pytest.mark.parametrize("extension", [".xls", ".xlsm", ".csv", ".xltx", ""])
def test_unsupported_format(tmp_path, extension):
    path = tmp_path / f"synthetic{extension}"
    path.write_bytes(b"not a workbook")
    with pytest.raises(UnsupportedExcelFormatError):
        ExcelReader(path)


def test_missing_and_directory(tmp_path):
    with pytest.raises(ExcelFileNotFoundError):
        ExcelReader(tmp_path / "missing.xlsx")
    with pytest.raises(ExcelWorkbookReadError):
        ExcelReader(tmp_path)


@pytest.mark.parametrize("content", [b"", b"not a zip", b"PK\x03\x04truncated"])
def test_corrupt_file_sanitized_and_closed(tmp_path, content):
    path = tmp_path / "synthetic-private-name.xlsx"
    path.write_bytes(content)
    with pytest.raises(ExcelWorkbookReadError) as caught:
        ExcelReader(path)
    assert str(path) not in str(caught.value)
    assert caught.value.__suppress_context__
    # Windows에서도 실패한 로딩의 handle이 남아 있지 않아야 한다.
    path.rename(tmp_path / "moved.xlsx")


def test_uppercase_extension(workbook_path):
    renamed = workbook_path.rename(workbook_path.with_suffix(".XLSX"))
    with ExcelReader(str(renamed)) as reader:
        assert reader.workbook_info.sheet_count == 3


def test_unknown_sheet(workbook_path):
    with ExcelReader(workbook_path) as reader:
        with pytest.raises(ExcelSheetNotFoundError):
            reader.sheet_info("missing")


@pytest.mark.parametrize("start,end", [(0, 1), (-1, 1), (2, 1), (1, 6), (True, 1), (1, 1.5)])
def test_invalid_header_range(workbook_path, start, end):
    with ExcelReader(workbook_path) as reader:
        with pytest.raises(InvalidHeaderRangeError):
            columns(reader, start=start, end=end)


def test_blank_and_duplicate_headers_preserved(workbook_path):
    with ExcelReader(workbook_path) as reader:
        result = columns(reader)
        assert len(result) == 8
        assert result[4].display_header == ""
        assert result[4].header_parts[0].cell.value is None
        assert result[6].display_header == result[7].display_header == "duplicate"
        assert [c.column_index for c in result] == list(range(1, 9))
        assert [c.column_letter for c in result] == list("ABCDEFGH")


def test_multirow_headers_and_true_merges(workbook_path):
    with ExcelReader(workbook_path) as reader:
        result = columns(reader, "MultiHeader", 2, 4)
        assert result[0].display_header == "Synthetic group | left | leaf"
        assert result[1].display_header == "Synthetic group | middle | 7"
        assert result[2].display_header == "Synthetic group | right | 2020-02-03 00:00:00"
        assert result[3].display_header == "Vertical"
        assert [p.cell.row_index for p in result[1].header_parts] == [2, 3, 4]
        merged = result[1].header_parts[0]
        assert merged.cell.coordinate == "B2" and merged.cell.value is None
        assert merged.merged_anchor.coordinate == "A2"
        assert merged.merged_range == "A2:C2"
        assert merged.effective_value == "Synthetic group"
        assert result[4].header_parts[0].effective_value is None
        assert result[4].header_parts[0].merged_anchor is None
        assert result[5].display_header == result[6].display_header == ""
        assert result[1].header_parts[2].cell.value == 7
        assert isinstance(result[2].header_parts[2].cell.value, datetime)
        # 헤더 선택이 병합 anchor보다 뒤에서 시작해도 실제 anchor를 추적한다.
        partial = columns(reader, "MultiHeader", 3, 4)[3]
        assert partial.header_parts[0].merged_anchor.coordinate == "D2"
        assert partial.display_header == "Vertical"


def test_native_values_types_and_provenance(workbook_path):
    with ExcelReader(workbook_path) as reader:
        result = list(rows(reader))
        assert [r.row_index for r in result] == [2, 3, 4, 5]
        cells = result[0].cells
        assert [c.column_index for c in cells] == list(range(1, 9))
        assert all(c.row_index == 2 for c in cells)
        assert cells[0].value == "001" and cells[0].value_type == "s"
        assert cells[1].value == 1 and type(cells[1].value) is int
        assert cells[1].number_format == "000" and cells[1].value_type == "n"
        assert cells[2].value == 3.5 and type(cells[2].value) is float
        assert cells[3].value == datetime(2020, 1, 2) and cells[3].value_type == "d"
        assert cells[4].value is None
        assert cells[5].is_formula and cells[5].value == "=B2+C2"
        assert cells[5].value_type == "f" and cells[5].coordinate == "F2"
        assert cells[6].value == "3.50" and not cells[6].is_formula
        assert cells[7].value == "#N/A" and cells[7].value_type == "e"
        assert result[1].cells[0].value is False and result[1].cells[0].value_type == "b"
        assert result[1].cells[1].value == 0
        assert result[1].cells[2].value == " "
        assert result[1].cells[3].value == datetime(2020, 1, 2, 3, 4, 5)
        assert result[1].cells[4].value == time(1, 2)


def test_blank_and_trailing_rows(workbook_path):
    with ExcelReader(workbook_path) as reader:
        result = list(rows(reader))
        assert [r.is_blank for r in result] == [False, False, True, False]
        filtered = list(rows(reader, include_blank=False))
        assert [r.row_index for r in filtered] == [2, 3, 5]
        assert all(c.value is None for c in filtered[-1].cells[:7])
        assert filtered[-1].cells[7].value == "synthetic trailing value"


@pytest.mark.parametrize("start", [0, -1, 1, 7, True, "2", 2.5])
def test_invalid_data_start(workbook_path, start):
    with ExcelReader(workbook_path) as reader:
        with pytest.raises(InvalidDataRangeError):
            rows(reader, data_start_row=start)


def test_custom_start_and_header_only(workbook_path):
    with ExcelReader(workbook_path) as reader:
        assert [r.row_index for r in rows(reader, data_start_row=5)] == [5]
        assert list(rows(reader, data_start_row=6)) == []
        assert list(reader.iter_rows("SingleHeader", header_start_row=1, header_end_row=5)) == []
        multi = list(
            reader.iter_rows("MultiHeader", header_start_row=2, header_end_row=4, data_start_row=6)
        )
        assert len(multi) == 1 and multi[0].cells[6].value == "synthetic last column"


@pytest.mark.parametrize("count,expected", [(0, []), (1, [2]), (2, [2, 3]), (50, [2, 3, 4, 5])])
def test_preview_bounds_metadata(workbook_path, count, expected):
    with ExcelReader(workbook_path) as reader:
        result = reader.read_preview(
            "SingleHeader", header_start_row=1, header_end_row=1, row_count=count
        )
        assert [r.row_index for r in result.rows] == expected
        assert result.columns == columns(reader)
        assert result.sheet.name == "SingleHeader"
        custom = reader.read_preview(
            "SingleHeader", header_start_row=1, header_end_row=1, start_row=4, row_count=1
        )
        assert len(custom.rows) == 1 and custom.rows[0].is_blank


@pytest.mark.parametrize(
    "options",
    [
        {"row_count": -1},
        {"row_count": True},
        {"row_count": 1.5},
        {"start_row": 0},
        {"start_row": 1},
        {"start_row": 7},
    ],
)
def test_invalid_preview(workbook_path, options):
    with ExcelReader(workbook_path) as reader:
        with pytest.raises(InvalidDataRangeError):
            reader.read_preview("SingleHeader", header_start_row=1, header_end_row=1, **options)


def test_source_unchanged_no_database_access_and_close(workbook_path, monkeypatch):
    before = hashlib.sha256(workbook_path.read_bytes()).digest()
    files_before = set(workbook_path.parent.iterdir())

    def forbidden(*args, **kwargs):
        pytest.fail("Reader must not save a workbook or open a database")

    monkeypatch.setattr(Workbook, "save", forbidden)
    monkeypatch.setattr(sqlite3, "connect", forbidden)
    reader = ExcelReader(workbook_path)
    with pytest.raises(RuntimeError), reader:
        columns(reader, "MultiHeader", 2, 4)
        list(rows(reader))
        reader.read_preview("SingleHeader", header_start_row=1, header_end_row=1)
        raise RuntimeError("synthetic caller failure")
    reader.close()
    with pytest.raises(ExcelReaderClosedError):
        _ = reader.workbook_info
    assert hashlib.sha256(workbook_path.read_bytes()).digest() == before
    assert set(workbook_path.parent.iterdir()) == files_before
    workbook_path.rename(workbook_path.with_name("moved.xlsx"))


def test_iterator_lazy_and_closed_guard(workbook_path, monkeypatch):
    import small_stream_research_tool.services.excel_reader as module

    original, calls = module._cell, []

    def observe(sheet, row, column):
        calls.append(row)
        return original(sheet, row, column)

    monkeypatch.setattr(module, "_cell", observe)
    reader = ExcelReader(workbook_path)
    iterator = rows(reader)
    assert iter(iterator) is iterator and calls == []
    assert next(iterator).row_index == 2 and set(calls) == {2}
    reader.close()
    with pytest.raises(ExcelReaderClosedError):
        next(iterator)


def test_open_failure_closes_owned_handle(workbook_path, monkeypatch):
    import small_stream_research_tool.services.excel_reader as module

    handles = []

    def fail(source, **kwargs):
        handles.append(source)
        raise ValueError("synthetic private content")

    monkeypatch.setattr(module, "load_workbook", fail)
    with pytest.raises(ExcelWorkbookReadError) as caught:
        ExcelReader(workbook_path)
    assert handles[0].closed
    assert "private" not in str(caught.value)


def test_special_native_values_and_chart_sheet(tmp_path):
    path = tmp_path / "special.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["array", "table", "duration", "error text"])
    sheet["A2"] = ArrayFormula(ref="A2:A3", text="=B2:B3*2")
    sheet["B2"] = DataTableFormula(ref="B2:B3", r1="D2")
    sheet["C2"] = timedelta(hours=27)
    sheet["D2"] = "#N/A"
    sheet["D2"].data_type = "s"
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=3, min_row=1, max_row=2), titles_from_data=True)
    chartsheet = workbook.create_chartsheet("Chart")
    chartsheet.add_chart(chart)
    workbook.save(path)
    workbook.close()
    with ExcelReader(path) as reader:
        cells = next(reader.iter_rows("Sheet", header_start_row=1, header_end_row=1)).cells
        assert isinstance(cells[0].value, ExcelFormula)
        assert cells[0].is_formula and cells[0].value.expression == "=B2:B3*2"
        assert dict(cells[0].value.attributes)["ref"] == "A2:A3"
        assert cells[1].is_formula and cells[1].value.kind == "dataTable"
        assert dict(cells[1].value.attributes)["r1"] == "D2"
        assert cells[2].value == timedelta(hours=27)
        assert cells[3].value == "#N/A" and cells[3].value_type == "s"
        assert reader.workbook_info.sheet_names == ("Sheet", "Chart")
        assert reader.sheet_info("Chart").max_row is None
        with pytest.raises(UnsupportedExcelSheetError):
            columns(reader, "Chart")
