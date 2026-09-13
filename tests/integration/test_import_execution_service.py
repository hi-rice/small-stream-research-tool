"""합성 Prepared 입력과 tmp SQLite로 A/B/C 경계 및 출처를 검증한다."""

import json
import sqlite3
import traceback
from contextlib import closing
from dataclasses import FrozenInstanceError, replace
from types import SimpleNamespace

import pytest
from openpyxl.utils import get_column_letter

from small_stream_research_tool.database import connect_database, initialize_database
from small_stream_research_tool.database.connection import transaction
from small_stream_research_tool.models.column_mapping import (
    ColumnMappingDraft,
    MappingMethod,
    MappingStatus,
)
from small_stream_research_tool.models.excel import ExcelCell, ExcelRow
from small_stream_research_tool.models.import_execution import (
    ImportExecutionRequest,
    ImportSheetSnapshot,
    SourceFileSnapshot,
)
from small_stream_research_tool.models.import_execution_errors import (
    DuplicateImportError,
    ImportFinalizeError,
    ImportPreflightError,
    ImportTransactionError,
)
from small_stream_research_tool.models.import_preparation import (
    ImportFieldPolicy,
    ImportPreparationResult,
    ImportPreparationSummary,
    PreparationStatus,
    PreparedCharacteristicValue,
    PreparedImportRow,
    PreparedStreamData,
    RowAction,
)
from small_stream_research_tool.models.import_preview import PreviewStatus
from small_stream_research_tool.repositories.dictionary_repository import DictionaryRepository
from small_stream_research_tool.repositories.stream_lookup_repository import StreamLookupRepository
from small_stream_research_tool.repositories.user_repository import UserRepository
from small_stream_research_tool.services.dictionary_service import DictionaryService
from small_stream_research_tool.services.import_execution_service import (
    FAILURE_CODE,
    FAILURE_MESSAGE,
    ImportExecutionService,
)
from small_stream_research_tool.services.import_preparation_service import ImportPreparationService
from small_stream_research_tool.services.import_preview_service import ImportPreviewService

STAMP = "2026-09-13T01:02:03Z"
POLICY = ImportFieldPolicy(frozenset({"synthetic_secret"}))
STREAM = "12345678009"
SECRET = "SYNTHETIC_PRIVATE_VALUE"
SOURCE = SourceFileSnapshot(
    "synthetic.xlsx", "C:/synthetic-only/synthetic.xlsx", ".xlsx", 12, "a" * 64, STAMP
)
SHEET = ImportSheetSnapshot("Synthetic Sheet", 1, 2, 3, 1)
TABLES = (
    "source_file",
    "import_history",
    "import_sheet",
    "import_column_mapping",
    "small_stream",
    "characteristic_value",
)


def core(code=STREAM):
    return PreparedStreamData(code, code[:2], code[2:5], code[5:8], code[8:], "Synthetic A")


def row(
    values=(),
    *,
    number=3,
    code=STREAM,
    action=RowAction.CREATE_STREAM,
    status=PreparationStatus.READY,
):
    return PreparedImportRow(
        number,
        code,
        "Synthetic Source Name",
        action,
        status,
        tuple(values),
        (),
        PreviewStatus.NEW_STREAM
        if action == RowAction.CREATE_STREAM
        else PreviewStatus.EXISTING_STREAM,
        core(code) if action == RowAction.CREATE_STREAM else None,
    )


def draft(item_id, column=4, status=MappingStatus.AUTO_MAPPED, method=MappingMethod.AUTO_ALIAS):
    return ColumnMappingDraft(
        column, get_column_letter(column), "Synthetic Header", (), item_id, method, status
    )


def preparation(rows):
    # 의도적으로 틀린 caller summary: 실행은 실제 rows/삽입으로 재계산해야 한다.
    return ImportPreparationResult(
        tuple(rows), ImportPreparationSummary(999, 999, 0, 0, 999, 0, 999)
    )


@pytest.fixture
def ctx(tmp_path):
    path = tmp_path / "synthetic.sqlite3"
    initialize_database(path)
    with closing(connect_database(path)) as conn:
        dictionary = DictionaryService(DictionaryRepository(conn))
        category = dictionary.create_category("synthetic", "Synthetic")
        version = dictionary.create_version("synthetic-v1")
        unit = dictionary.create_unit("Synthetic Unit", "syn")
        item = dictionary.create_item(
            "Synthetic Metric",
            "synthetic_metric",
            category.category_id,
            "REAL",
            version.version_id,
            unit_id=unit.unit_id,
        )
        with transaction(conn):
            user = UserRepository(conn).create_user(
                login_id="synthetic_user",
                password_hash="SYNTHETIC_TEST_HASH_ONLY",
                display_name="Synthetic User",
                department=None,
                role=None,
                timestamp=STAMP,
            )
        value = PreparedCharacteristicValue(
            item.dictionary_id,
            "REAL",
            4,
            3,
            value_number=1.25,
            original_value=" 1.250 ",
            original_unit="syn",
            unit_id=unit.unit_id,
        )
        request = ImportExecutionRequest(
            SOURCE,
            SHEET,
            preparation([row([value])]),
            (draft(item.dictionary_id),),
            user.user_id,
            version.version_id,
            "synthetic-batch",
            2026,
        )
        yield SimpleNamespace(
            conn=conn,
            path=path,
            dictionary=dictionary,
            category=category,
            version=version,
            unit=unit,
            item=item,
            value=value,
            user=user,
            request=request,
            service=ImportExecutionService(conn),
        )


def execute(ctx, request=None, policy=POLICY):
    return ctx.service.execute(request or ctx.request, field_policy=policy)


def counts(ctx):
    return tuple(
        ctx.conn.execute(f"SELECT count(*) FROM {table}").fetchone()[0] for table in TABLES
    )


def assert_no_writes(ctx, request, policy=POLICY, error=ImportPreflightError):
    before = list(ctx.conn.iterdump())
    writes = []

    def authorize(action, arg1, _arg2, _db, _trigger):
        if action in (sqlite3.SQLITE_INSERT, sqlite3.SQLITE_UPDATE, sqlite3.SQLITE_DELETE):
            writes.append(arg1)
            return sqlite3.SQLITE_DENY
        return sqlite3.SQLITE_OK

    ctx.conn.set_authorizer(authorize)
    try:
        with pytest.raises(error):
            execute(ctx, request, policy)
    finally:
        ctx.conn.set_authorizer(None)
    assert writes == []
    assert list(ctx.conn.iterdump()) == before


def test_success_provenance_defaults_and_counts(ctx):
    result = execute(ctx)
    assert counts(ctx) == (1, 1, 1, 1, 1, 1)
    assert result.status == "SUCCESS" and result.data_committed
    assert (
        result.ready_rows,
        result.excluded_rows,
        result.created_stream_count,
        result.reused_stream_count,
        result.characteristic_value_count,
    ) == (1, 0, 1, 0, 1)
    history = ctx.service._histories.get_by_id(result.import_id)
    assert history.created_by_user_id == ctx.user.user_id
    assert history.dictionary_version_id == ctx.version.version_id
    assert (
        history.schema_version_id
        == ctx.conn.execute("SELECT schema_version_id FROM schema_version").fetchone()[0]
    )
    assert (
        history.total_rows,
        history.accepted_rows,
        history.warning_rows,
        history.rejected_rows,
    ) == (1, 1, None, 0)
    assert history.started_at.endswith("Z") and history.finished_at.endswith("Z")
    value = ctx.service._values.list_by_import_id(result.import_id)[0]
    provenance = ctx.service._values.get_provenance(value.characteristic_value_id)
    assert provenance.history.import_id == result.import_id
    assert provenance.source_file.source_file_id == result.source_file_id
    assert provenance.sheet.import_id == result.import_id
    assert provenance.mapping.import_sheet_id == value.import_sheet_id
    assert provenance.mapping.dictionary_id == value.dictionary_id
    assert provenance.mapping.source_column_index == 4 and value.source_row == 3
    assert value.value_number == 1.25 and value.original_value == " 1.250 "
    assert value.original_unit == "syn" and value.unit_id == ctx.unit.unit_id
    assert value.reference_year == 2026
    assert not value.is_representative and value.is_active and value.quality_status == "UNREVIEWED"
    assert value.source_type == "IMPORT"
    with pytest.raises(FrozenInstanceError):
        result.status = "FAILED"


@pytest.mark.parametrize(
    "mutation",
    [
        "blocked",
        "all_excluded",
        "empty",
        "unknown_status",
        "bad_user",
        "bool_user",
        "bad_version",
        "bad_batch",
        "empty_name",
        "empty_path",
        "bad_size",
        "bad_sheet",
        "header_order",
        "data_order",
        "row_zero",
        "duplicate_row",
        "duplicate_column",
        "missing_mapping",
        "wrong_dictionary",
        "unmapped_value",
        "wrong_source_row",
        "wrong_unit",
        "missing_core",
        "core_mismatch",
        "invalid_action",
        "duplicate_value",
        "bad_year",
        "unknown_dictionary",
        "invalid_code",
        "existing_empty",
    ],
)
def test_invalid_preflight_zero_writes(ctx, mutation):
    req = ctx.request
    current = req.preparation.rows[0]
    row_changes = {
        "blocked": dict(status=PreparationStatus.BLOCKED),
        "all_excluded": dict(status=PreparationStatus.EXCLUDED),
        "unknown_status": dict(status="READY"),
        "row_zero": dict(source_row=0),
        "wrong_source_row": dict(prepared_values=(replace(ctx.value, source_row=4),)),
        "wrong_unit": dict(prepared_values=(replace(ctx.value, unit_id=None),)),
        "missing_core": dict(core_data=None),
        "core_mismatch": dict(core_data=core("12345678008")),
        "invalid_action": dict(row_action=None),
        "duplicate_value": dict(prepared_values=(ctx.value, ctx.value)),
        "invalid_code": dict(stream_code="123"),
        "existing_empty": dict(
            row_action=RowAction.USE_EXISTING_STREAM, core_data=None, prepared_values=()
        ),
    }
    if mutation in row_changes:
        req = replace(req, preparation=preparation([replace(current, **row_changes[mutation])]))
    elif mutation == "empty":
        req = replace(req, preparation=preparation([]))
    elif mutation == "duplicate_row":
        req = replace(req, preparation=preparation([current, current]))
    elif mutation in ("bad_user", "bool_user", "bad_version", "bad_batch", "bad_year"):
        key, value = {
            "bad_user": ("current_user_id", 999),
            "bool_user": ("current_user_id", True),
            "bad_version": ("dictionary_version_id", 999),
            "bad_batch": ("batch_code", " "),
            "bad_year": ("reference_year", True),
        }[mutation]
        req = replace(req, **{key: value})
    elif mutation in ("empty_name", "empty_path", "bad_size"):
        key, value = {
            "empty_name": ("file_name", ""),
            "empty_path": ("original_path", ""),
            "bad_size": ("file_size", -1),
        }[mutation]
        req = replace(req, source=replace(SOURCE, **{key: value}))
    elif mutation in ("bad_sheet", "header_order", "data_order"):
        key, value = {
            "bad_sheet": ("sheet_name", ""),
            "header_order": ("header_start_row", 3),
            "data_order": ("data_start_row", 2),
        }[mutation]
        req = replace(req, sheet=replace(SHEET, **{key: value}))
    else:
        mappings = {
            "duplicate_column": req.column_mappings * 2,
            "missing_mapping": (),
            "wrong_dictionary": (draft(ctx.item.dictionary_id, 5),),
            "unknown_dictionary": (draft(999),),
            "unmapped_value": (
                draft(None, status=MappingStatus.UNMAPPED, method=MappingMethod.NONE),
            ),
        }
        req = replace(req, column_mappings=mappings[mutation])
    assert_no_writes(ctx, req)


def test_explicit_policy_required(ctx):
    assert_no_writes(ctx, ctx.request, None)
    with pytest.raises(ImportPreflightError):
        ctx.service.execute(ctx.request)
    assert counts(ctx) == (0,) * 6


def test_sensitive_policy_rechecked_without_header_heuristics(ctx):
    assert_no_writes(ctx, ctx.request, ImportFieldPolicy(frozenset({"synthetic_metric"})))
    req = replace(
        ctx.request,
        column_mappings=(
            replace(ctx.request.column_mappings[0], source_header="password header example"),
        ),
    )
    assert execute(ctx, req).status == "SUCCESS"


@pytest.mark.parametrize("name", ["stream_code", "stream_name", "province_code"])
def test_core_policy_recheck(ctx, name):
    assert_no_writes(ctx, ctx.request, ImportFieldPolicy(frozenset({name})))


def test_excluded_payload_never_written(ctx):
    excluded = row(
        [replace(ctx.value, source_row=4, original_value=SECRET)],
        number=4,
        code="12345678008",
        status=PreparationStatus.EXCLUDED,
    )
    result = execute(
        ctx,
        replace(ctx.request, preparation=preparation([ctx.request.preparation.rows[0], excluded])),
    )
    assert result.excluded_rows == 1 and result.characteristic_value_count == 1
    assert SECRET not in "\n".join(ctx.conn.iterdump())
    assert ctx.service._streams.count() == 1
    history = ctx.service._histories.get_by_id(result.import_id)
    assert (history.total_rows, history.accepted_rows, history.rejected_rows) == (2, 1, 0)


@pytest.mark.parametrize(
    "table, column",
    [
        ("app_user", "is_active"),
        ("data_dictionary", "is_active"),
        ("data_category", "is_active"),
        ("unit_dictionary", "is_active"),
    ],
)
def test_inactive_reference_zero_writes(ctx, table, column):
    ctx.conn.execute(f"UPDATE {table} SET {column}=0")
    assert_no_writes(ctx, ctx.request)


def test_caller_transaction_preserved(ctx):
    with transaction(ctx.conn):
        assert_no_writes(ctx, ctx.request)
        assert ctx.conn.in_transaction


def test_source_and_running_are_visible_before_b(ctx, monkeypatch):
    original = ctx.service._transaction_b

    def inspect(*args):
        with closing(connect_database(ctx.path)) as observer:
            assert observer.execute("SELECT status FROM import_history").fetchone()[0] == "RUNNING"
            assert observer.execute("SELECT count(*) FROM source_file").fetchone()[0] == 1
            assert observer.execute("SELECT count(*) FROM import_sheet").fetchone()[0] == 0
        return original(*args)

    monkeypatch.setattr(ctx.service, "_transaction_b", inspect)
    execute(ctx)


def fail(*_args, **_kwargs):
    raise sqlite3.OperationalError(SECRET + SOURCE.original_path)


def test_a_history_failure_rolls_back_source(ctx, monkeypatch):
    monkeypatch.setattr(ctx.service._histories, "create", fail)
    with pytest.raises(ImportTransactionError) as error:
        execute(ctx)
    assert error.value.phase == "A" and error.value.result.import_id is None
    assert counts(ctx) == (0,) * 6


def test_a_source_failure_no_b_or_c(ctx, monkeypatch):
    monkeypatch.setattr(ctx.service._sources, "create", fail)
    monkeypatch.setattr(ctx.service, "_transaction_b", lambda *_: pytest.fail("B called"))
    monkeypatch.setattr(ctx.service, "_finalize", lambda *_: pytest.fail("C called"))
    with pytest.raises(ImportTransactionError) as error:
        execute(ctx)
    assert error.value.phase == "A" and counts(ctx) == (0,) * 6


def test_same_hash_new_source_each_run_and_generated_batch(ctx):
    first = execute(ctx, replace(ctx.request, batch_code=None))
    existing = row([ctx.value], action=RowAction.USE_EXISTING_STREAM)
    second = execute(
        ctx, replace(ctx.request, batch_code=None, preparation=preparation([existing]))
    )
    assert first.batch_code != second.batch_code
    assert first.source_file_id != second.source_file_id
    assert len(ctx.service._sources.find_by_hash(SOURCE.file_hash)) == 2
    assert second.reused_stream_count == 1


@pytest.mark.parametrize(
    "status, method, stored",
    [
        (MappingStatus.AUTO_MAPPED, MappingMethod.AUTO_ALIAS, "AUTO"),
        (MappingStatus.USER_MAPPED, MappingMethod.USER, "USER"),
        (MappingStatus.UNMAPPED, MappingMethod.NONE, "NONE"),
        (MappingStatus.DO_NOT_MAP, MappingMethod.USER, "NONE"),
        (MappingStatus.NEEDS_REVIEW, MappingMethod.AUTO_ALIAS, "AUTO_ALIAS"),
        (MappingStatus.NEEDS_REVIEW, MappingMethod.USER, "USER"),
    ],
)
def test_mapping_adapter_preserves_status(ctx, status, method, stored):
    dictionary_id = (
        None
        if status in (MappingStatus.UNMAPPED, MappingStatus.DO_NOT_MAP)
        else ctx.item.dictionary_id
    )
    req = replace(
        ctx.request,
        preparation=preparation([row()]),
        column_mappings=(draft(dictionary_id, status=status, method=method),),
    )
    result = execute(ctx, req)
    mapping = ctx.service._mappings.list_by_import_sheet_id(
        ctx.service._sheets.list_by_import_id(result.import_id)[0].import_sheet_id
    )[0]
    assert mapping.mapping_status == status and mapping.mapping_method == stored
    assert mapping.dictionary_id == dictionary_id
    assert result.characteristic_value_count == 0 and result.created_stream_count == 1


def test_create_conflict_b_rollback(ctx):
    with transaction(ctx.conn):
        ctx.service._streams.create_prepared(core(), timestamp=STAMP)
    with pytest.raises(ImportTransactionError) as error:
        execute(ctx)
    assert error.value.phase == "B"
    assert counts(ctx) == (1, 1, 0, 0, 1, 0)
    assert ctx.service._streams.get_by_stream_code(STREAM).created_at == STAMP


def test_missing_existing_b_rollback(ctx):
    req = replace(
        ctx.request,
        preparation=preparation([row([ctx.value], action=RowAction.USE_EXISTING_STREAM)]),
    )
    with pytest.raises(ImportTransactionError) as error:
        execute(ctx, req)
    assert error.value.phase == "B" and counts(ctx) == (1, 1, 0, 0, 0, 0)


def test_existing_core_and_representative_untouched(ctx):
    with transaction(ctx.conn):
        stream = ctx.service._streams.create_prepared(core(), timestamp=STAMP)
        representative = ctx.service._values.create(
            stream_code=STREAM,
            dictionary_id=ctx.item.dictionary_id,
            value_number=9.0,
            is_representative=1,
            created_at=STAMP,
            updated_at=STAMP,
        )
        ctx.conn.execute(
            "INSERT INTO stream_characteristic VALUES (?, ?, ?, ?)",
            (STREAM, ctx.item.dictionary_id, representative.characteristic_value_id, STAMP),
        )
    untouched = {
        name: ctx.conn.execute(f"SELECT * FROM {name}").fetchall()
        for name in ("stream_characteristic", "data_quality_issue", "record_history")
    }
    req = replace(
        ctx.request,
        preparation=preparation([row([ctx.value], action=RowAction.USE_EXISTING_STREAM)]),
    )
    result = execute(ctx, req)
    assert result.reused_stream_count == 1 and result.created_stream_count == 0
    assert ctx.service._streams.get_by_stream_code(STREAM) == stream
    assert ctx.service._values.get_by_id(representative.characteristic_value_id) == representative
    for name, records in untouched.items():
        assert ctx.conn.execute(f"SELECT * FROM {name}").fetchall() == records


@pytest.mark.parametrize(
    "data_type, field, value",
    [
        ("REAL", "value_number", 1.23456789),
        ("INTEGER", "value_integer", 123),
        ("TEXT", "value_text", " Synthetic text "),
        ("DATE", "value_date", "2026-01-02"),
        ("DATETIME", "value_date", "2026-01-02T03:04:05+09:00"),
    ],
)
def test_typed_values_not_reparsed_or_converted(ctx, data_type, field, value, monkeypatch):
    item = ctx.dictionary.create_item(
        "Synthetic Other",
        "synthetic_other",
        ctx.category.category_id,
        data_type,
        ctx.version.version_id,
    )
    prepared = PreparedCharacteristicValue(
        item.dictionary_id, data_type, 7, 3, **{field: value}, original_value=" Synthetic original "
    )
    monkeypatch.setattr(
        "small_stream_research_tool.services.value_normalization_service.normalize_cell", fail
    )
    req = replace(
        ctx.request,
        preparation=preparation([row([prepared])]),
        column_mappings=(draft(item.dictionary_id, 7),),
    )
    result = execute(ctx, req)
    stored = ctx.service._values.list_by_import_id(result.import_id)[0]
    assert getattr(stored, field) == value and stored.original_value == prepared.original_value


@pytest.mark.parametrize("repository", ["_sheets", "_mappings", "_streams", "_values"])
def test_b_failure_all_data_rollback_and_safe_c(ctx, monkeypatch, repository, caplog):
    repo = getattr(ctx.service, repository)
    monkeypatch.setattr(repo, "create", fail)
    with pytest.raises(ImportTransactionError) as error:
        execute(ctx)
    result = error.value.result
    assert error.value.phase == "B" and result.status == "FAILED"
    assert not result.data_committed
    assert (
        result.created_stream_count,
        result.reused_stream_count,
        result.characteristic_value_count,
    ) == (0, 0, 0)
    assert counts(ctx) == (1, 1, 0, 0, 0, 0)
    history = ctx.service._histories.get_by_id(result.import_id)
    assert history.status == "FAILED" and history.finished_at.endswith("Z")
    assert history.error_code == FAILURE_CODE and history.error_message == FAILURE_MESSAGE
    assert history.accepted_rows == 0 and history.rejected_rows is None
    visible = "".join(traceback.format_exception(error.value)) + caplog.text + repr(result)
    assert SECRET not in visible and SOURCE.original_path not in visible
    assert SECRET not in "\n".join(ctx.conn.iterdump())


def test_last_row_failure_rolls_back_earlier_rows(ctx, monkeypatch):
    second = row([replace(ctx.value, source_row=4)], number=4, code="12345678008")
    req = replace(ctx.request, preparation=preparation([ctx.request.preparation.rows[0], second]))
    original = ctx.service._values.create_prepared
    calls = []

    def insert(prepared, **kwargs):
        calls.append(prepared.source_row)
        if len(calls) == 2:
            fail()
        return original(prepared, **kwargs)

    monkeypatch.setattr(ctx.service._values, "create_prepared", insert)
    with pytest.raises(ImportTransactionError):
        execute(ctx, req)
    assert calls == [3, 4] and counts(ctx) == (1, 1, 0, 0, 0, 0)


@pytest.mark.parametrize("mismatch", ["import_sheet_id", "dictionary_id", "source_column_index"])
def test_mapping_reference_mismatch_rolls_back(ctx, monkeypatch, mismatch):
    original = ctx.service._mappings.create

    def create(**kwargs):
        record = original(**kwargs)
        return replace(record, **{mismatch: 999})

    monkeypatch.setattr(ctx.service._mappings, "create", create)
    with pytest.raises(ImportTransactionError):
        execute(ctx)
    assert counts(ctx) == (1, 1, 0, 0, 0, 0)


@pytest.mark.parametrize("b_failed", [False, True])
def test_finalize_failure_no_retry_or_repair(ctx, monkeypatch, b_failed):
    if b_failed:
        monkeypatch.setattr(ctx.service._values, "create", fail)
    original = ctx.service._histories.update_status
    calls = []

    def update(*args, **kwargs):
        calls.append(1)
        original(*args, **kwargs)
        fail()  # C UPDATE 후 실패도 C만 rollback.

    monkeypatch.setattr(ctx.service._histories, "update_status", update)
    with pytest.raises(ImportFinalizeError) as error:
        execute(ctx)
    result = error.value.result
    assert error.value.recovery_required and error.value.phase == "C"
    assert result.status == "RUNNING" and result.finished_at is None
    assert result.data_committed is not b_failed
    assert result.characteristic_value_count == (0 if b_failed else 1)
    assert counts(ctx) == ((1, 1, 0, 0, 0, 0) if b_failed else (1, 1, 1, 1, 1, 1))
    history = ctx.service._histories.get_by_id(result.import_id)
    assert history.status == "RUNNING" and history.finished_at is None
    assert history.accepted_rows is None
    assert_no_writes(ctx, ctx.request, error=DuplicateImportError)
    assert calls == [1]


def test_success_duplicate_batch_zero_writes(ctx):
    execute(ctx)
    assert_no_writes(ctx, ctx.request, error=DuplicateImportError)


def test_settings_whitelist_and_repr(ctx):
    req = replace(
        ctx.request, preparation=preparation([row([replace(ctx.value, original_value=SECRET)])])
    )
    result = execute(ctx, req)
    settings = ctx.service._histories.get_by_id(result.import_id).settings_json
    expected = dict(
        sheet_name=SHEET.sheet_name,
        header_start_row=1,
        header_end_row=2,
        data_start_row=3,
        dictionary_version_id=ctx.version.version_id,
    )
    assert settings == json.dumps(expected, sort_keys=True, ensure_ascii=True)
    for text in (settings, repr(req), repr(result), repr(SOURCE)):
        assert SECRET not in text and SOURCE.original_path not in text


def test_schema_pk_not_assumed_and_version_not_guessed(ctx):
    ctx.conn.execute("UPDATE schema_version SET schema_version_id=71")
    result = execute(ctx, replace(ctx.request, dictionary_version_id=None))
    history = ctx.service._histories.get_by_id(result.import_id)
    assert history.schema_version_id == 71 and history.dictionary_version_id is None


def test_only_expected_tables_written_no_delete_or_replace(ctx):
    writes = []

    def authorize(action, arg1, _arg2, _db, _trigger):
        if action in (sqlite3.SQLITE_INSERT, sqlite3.SQLITE_UPDATE, sqlite3.SQLITE_DELETE):
            writes.append((action, arg1))
        return sqlite3.SQLITE_OK

    ctx.conn.set_authorizer(authorize)
    try:
        execute(ctx)
    finally:
        ctx.conn.set_authorizer(None)
    assert {table for _, table in writes} == set(TABLES)
    assert not any(action == sqlite3.SQLITE_DELETE for action, _ in writes)
    assert {table for action, table in writes if action == sqlite3.SQLITE_UPDATE} == {
        "import_history"
    }


def test_real_preview_preparation_execution_chain(ctx):
    definitions = (
        ("stream_code", STREAM),
        ("province_code", "12"),
        ("city_county_code", "345"),
        ("town_code", "678"),
        ("stream_serial_no", "009"),
        ("stream_name", "Synthetic A"),
    )
    mappings = []
    cells = []
    for column, (name, value) in enumerate(definitions, 1):
        item = ctx.dictionary.create_item(
            "Synthetic " + name, name, ctx.category.category_id, "TEXT", ctx.version.version_id
        )
        mappings.append(
            draft(item.dictionary_id, column, MappingStatus.USER_MAPPED, MappingMethod.USER)
        )
        cells.append(ExcelCell(3, column, get_column_letter(column), value, "s", False, "General"))
    mappings.append(draft(ctx.item.dictionary_id, 7, MappingStatus.USER_MAPPED, MappingMethod.USER))
    cells.append(ExcelCell(3, 7, "G", " 1.250 ", "s", False, "General"))
    preview = ImportPreviewService(ctx.dictionary, StreamLookupRepository(ctx.conn)).build_preview(
        [ExcelRow(3, tuple(cells))], tuple(mappings)
    )
    prepared = ImportPreparationService(ctx.dictionary).prepare(preview.rows, field_policy=POLICY)
    assert prepared.rows[0].status == PreparationStatus.READY
    result = execute(
        ctx, replace(ctx.request, preparation=prepared, column_mappings=tuple(mappings))
    )
    assert result.created_stream_count == result.characteristic_value_count == 1
    assert counts(ctx) == (1, 1, 1, 7, 1, 1)
    stored = ctx.service._values.list_by_import_id(result.import_id)[0]
    assert ctx.service._mappings.get_by_id(stored.mapping_id).source_column_index == 7
    assert stored.original_value == " 1.250 " and stored.value_number == 1.25


@pytest.mark.parametrize("phase", ["A", "B", "C"])
def test_actual_commit_denial_preserves_transaction_boundaries(ctx, phase):
    # cached COMMIT문의 authorizer는 prepare 때만 호출되므로, 매 단계 seam에서 재설정한다.
    original = getattr(
        ctx.service, {"A": "_transaction_a", "B": "_transaction_b", "C": "_finalize"}[phase]
    )

    def deny_commit(action, arg1, _arg2, _db, _trigger):
        if action == sqlite3.SQLITE_TRANSACTION and arg1 == "COMMIT":
            return sqlite3.SQLITE_DENY
        return sqlite3.SQLITE_OK

    def wrapped(*args, **kwargs):
        ctx.conn.set_authorizer(deny_commit)
        try:
            return original(*args, **kwargs)
        finally:
            ctx.conn.set_authorizer(None)

    name = {"A": "_transaction_a", "B": "_transaction_b", "C": "_finalize"}[phase]
    setattr(ctx.service, name, wrapped)
    error_type = ImportFinalizeError if phase == "C" else ImportTransactionError
    with pytest.raises(error_type) as error:
        execute(ctx)
    assert error.value.phase == phase
    assert not ctx.conn.in_transaction
    assert (
        counts(ctx)
        == {"A": (0, 0, 0, 0, 0, 0), "B": (1, 1, 0, 0, 0, 0), "C": (1, 1, 1, 1, 1, 1)}[phase]
    )
    if phase != "A":
        history = ctx.service._histories.get_by_id(error.value.result.import_id)
        assert history.status == ("RUNNING" if phase == "C" else "FAILED")


def test_blank_header_unmapped_metadata(ctx):
    unmapped = replace(draft(None, 8, MappingStatus.UNMAPPED, MappingMethod.NONE), source_header="")
    result = execute(
        ctx, replace(ctx.request, column_mappings=(*ctx.request.column_mappings, unmapped))
    )
    sheet = ctx.service._sheets.list_by_import_id(result.import_id)[0]
    records = ctx.service._mappings.list_by_import_sheet_id(sheet.import_sheet_id)
    assert records[-1].source_header == "" and records[-1].normalized_header is None
    assert result.characteristic_value_count == 1


@pytest.mark.parametrize("field", ["stream_code", "created_at"])
def test_core_or_system_disguised_as_characteristic_rejected(ctx, field):
    ctx.conn.execute("UPDATE data_dictionary SET internal_name=?", (field,))
    assert_no_writes(ctx, ctx.request)


@pytest.mark.parametrize("column, value", [("data_type", "INTEGER"), ("deprecated_version_id", 1)])
def test_changed_dictionary_rejected(ctx, column, value):
    ctx.conn.execute(f"UPDATE data_dictionary SET {column}=?", (value,))
    assert_no_writes(ctx, ctx.request)


def test_reference_changed_after_a_rechecked_inside_b(ctx, monkeypatch):
    original = ctx.service._transaction_a

    def create(*args):
        result = original(*args)
        ctx.conn.execute("UPDATE app_user SET is_active=0")
        return result

    monkeypatch.setattr(ctx.service, "_transaction_a", create)
    with pytest.raises(ImportTransactionError) as error:
        execute(ctx)
    assert error.value.phase == "B" and counts(ctx) == (1, 1, 0, 0, 0, 0)


def test_excluded_actual_sensitive_dictionary_not_written(ctx):
    item = ctx.dictionary.create_item(
        "Synthetic Secret",
        "synthetic_secret",
        ctx.category.category_id,
        "TEXT",
        ctx.version.version_id,
    )
    value = PreparedCharacteristicValue(
        item.dictionary_id, "TEXT", 8, 4, value_text=SECRET, original_value=SECRET
    )
    excluded = row([value], number=4, code="12345678008", status=PreparationStatus.EXCLUDED)
    req = replace(
        ctx.request,
        preparation=preparation([ctx.request.preparation.rows[0], excluded]),
        column_mappings=(*ctx.request.column_mappings, draft(item.dictionary_id, 8)),
    )
    result = execute(ctx, req)
    assert result.excluded_rows == 1 and SECRET not in "\n".join(ctx.conn.iterdump())


def test_multiple_rows_create_reuse_counts_and_column_lookup(ctx):
    with transaction(ctx.conn):
        ctx.service._streams.create_prepared(core("12345678008"), timestamp=STAMP)
    other = ctx.dictionary.create_item(
        "Synthetic Other",
        "synthetic_other",
        ctx.category.category_id,
        "TEXT",
        ctx.version.version_id,
    )
    second_value = PreparedCharacteristicValue(
        other.dictionary_id, "TEXT", 9, 3, value_text="Synthetic"
    )
    first = row([second_value, ctx.value])
    second = row(
        [replace(ctx.value, source_row=4)],
        number=4,
        code="12345678008",
        action=RowAction.USE_EXISTING_STREAM,
    )
    result = execute(
        ctx,
        replace(
            ctx.request,
            preparation=preparation([first, second]),
            column_mappings=(draft(other.dictionary_id, 9), *ctx.request.column_mappings),
        ),
    )
    assert (
        result.ready_rows,
        result.created_stream_count,
        result.reused_stream_count,
        result.characteristic_value_count,
    ) == (2, 1, 1, 3)
    for stored in ctx.service._values.list_by_import_id(result.import_id):
        mapping = ctx.service._mappings.get_by_id(stored.mapping_id)
        assert mapping.dictionary_id == stored.dictionary_id
        assert mapping.source_column_index == (
            9 if stored.dictionary_id == other.dictionary_id else 4
        )


def test_sheet_index_matches_reader_one_based(ctx):
    assert_no_writes(ctx, replace(ctx.request, sheet=replace(SHEET, sheet_index=0)))


def test_failed_batch_reentry_zero_writes(ctx, monkeypatch):
    monkeypatch.setattr(ctx.service._values, "create", fail)
    with pytest.raises(ImportTransactionError):
        execute(ctx)
    assert_no_writes(ctx, ctx.request, error=DuplicateImportError)


def test_source_optional_snapshot_fields_preserved(ctx):
    req = replace(ctx.request, source=SourceFileSnapshot("synthetic.xlsx", "synthetic.xlsx"))
    result = execute(ctx, req)
    source = ctx.service._sources.get_by_id(result.source_file_id)
    assert source.file_hash is None and source.file_size is None and source.file_modified_at is None
